"""Folder-first local document workspace with bounded, read-only previews.

Only a directory the user selected is enumerated. PDF rendering stays on the GUI
thread and is suspended before background jobs start. Text editing writes copies.
"""
from __future__ import annotations

import math
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from PySide6.QtCore import QFileInfo, QSize, Qt, QTimer, QUrl, Signal
from PySide6.QtGui import QDesktopServices, QImage, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView, QCheckBox, QFileDialog, QFileIconProvider, QFrame, QHBoxLayout,
    QInputDialog, QLabel, QLineEdit, QListWidget, QListWidgetItem, QMessageBox,
    QPlainTextEdit, QPushButton, QScrollArea, QSpinBox, QSplitter, QStackedWidget,
    QVBoxLayout, QWidget,
)

from .convert import atomic_save, fitz
from .util import IMAGE_EXTS, TEXT_EXTS, UserError, file_size, human_size, sibling

OFFICE_EXTS = {".docx", ".xlsx", ".pptx", ".odt"}
DOCUMENT_EXTS = OFFICE_EXTS | {".pdf", ".svg", ".epub", ".xps", ".oxps", ".fb2", ".cbz"}
MAX_FOLDER_ENTRIES = 5000
MAX_TEXT_BYTES = 1_048_576
PREVIEW_TEXT_LIMIT = 160_000
PREPARE_EDIT_ACTION = "Открыть копию в редакторе"


def _label(text, name="", wrap=False):
    widget = QLabel(text)
    widget.setObjectName(name)
    widget.setTextFormat(Qt.PlainText)
    widget.setWordWrap(wrap)
    return widget


def _button(text, callback, name=""):
    widget = QPushButton(text)
    widget.setObjectName(name)
    widget.setCursor(Qt.PointingHandCursor)
    widget.clicked.connect(callback)
    return widget


def _safe_xml(archive, name):
    info = archive.getinfo(name)
    if info.file_size > 16 * 1024 * 1024:
        raise UserError("Содержимое документа слишком велико для быстрого просмотра.")
    return ElementTree.fromstring(archive.read(name))


def office_text_preview(path: str | Path) -> str:
    """Bounded textual Office preview; never follows links, evaluates macros or formulas."""
    path = Path(path)
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 20000 or sum(info.file_size for info in infos) > 256 * 1024 * 1024:
            raise UserError("Документ слишком велик для быстрого просмотра.")
        suffix = path.suffix.lower()
        lines = []
        if suffix == ".docx":
            root = _safe_xml(archive, "word/document.xml")
            namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            for paragraph in root.findall(".//w:p", namespace):
                lines.append("".join(node.text or "" for node in paragraph.findall(".//w:t", namespace)))
        elif suffix == ".pptx":
            names = sorted((name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")), key=lambda name: int(Path(name).stem.removeprefix("slide")))
            for index, name in enumerate(names[:100]):
                root = _safe_xml(archive, name)
                texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
                lines.append(f"СЛАЙД {index + 1}\n" + "\n".join(texts))
        elif suffix == ".odt":
            root = _safe_xml(archive, "content.xml")
            lines = ["".join(node.itertext()) for node in root.iter() if node.tag.endswith(("}p", "}h"))]
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook.worksheets[:8]:
                    lines.append(f"ЛИСТ: {sheet.title}")
                    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 40, 40), max_col=min(sheet.max_column or 12, 12), values_only=True):
                        lines.append("\t".join("" if value is None else str(value) for value in row))
                    lines.append("")
            finally:
                workbook.close()
        return "\n".join(lines)[:PREVIEW_TEXT_LIMIT] or "В документе нет доступного текстового содержимого."


def read_text_preview(path: str | Path):
    path = Path(path)
    with path.open("rb") as source:
        data = source.read(MAX_TEXT_BYTES + 1)
    complete = len(data) <= MAX_TEXT_BYTES
    data = data[:MAX_TEXT_BYTES]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        encoding = "utf-8-sig"
    try:
        text = data.decode(encoding)
    except UnicodeError:
        encoding = "cp1251"
        text = data.decode(encoding, errors="replace")
    if "\x00" in text:
        raise UserError("Файл содержит двоичные данные; текстовый просмотр недоступен.")
    return text[:PREVIEW_TEXT_LIMIT], complete and len(text) <= PREVIEW_TEXT_LIMIT, encoding


class LibraryWorkspace(QWidget):
    tools_requested = Signal(list)
    convert_requested = Signal(list)
    file_selected = Signal(str)

    def __init__(self, controller, parent=None, store=None):
        super().__init__(parent)
        self.controller = controller
        self.store = store
        self.folder = None
        self.current_path = None
        self._pdf = None
        self._password = ""
        self._processing = False
        self._selecting = False
        self._entries = []
        self._text_editing = False
        self._text_complete = False
        self._pending_edit_source = None
        self._pending_editor = None
        self._preview_kind = "empty"
        self._icons = QFileIconProvider()
        self.setAcceptDrops(True)
        col = QVBoxLayout(self)
        col.setContentsMargins(25, 24, 25, 20)
        col.setSpacing(14)
        header = QHBoxLayout()
        titles = QVBoxLayout()
        titles.setSpacing(5)
        titles.addWidget(_label("DOCUFREEDOM  /  DNA.APPS", "eyebrow"))
        titles.addWidget(_label("Ваши документы — под рукой", "pageTitle"))
        header.addLayout(titles, 1)
        self.open_files_button = _button("Открыть файл…", self.choose_files)
        self.choose_folder_button = _button("Открыть папку…", self.choose_folder, "primary")
        header.addWidget(self.open_files_button)
        header.addWidget(self.choose_folder_button)
        col.addLayout(header)
        col.addWidget(_label("Выберите файл слева, просмотрите его здесь и откройте редактор одним нажатием.", "muted", True))
        paths = QHBoxLayout()
        self.up_button = _button("↑", self.parent_folder)
        self.up_button.setToolTip("В родительскую папку")
        paths.addWidget(self.up_button)
        self.path_button = _button("Выберите папку", self.choose_folder, "quiet")
        self.path_button.setSizePolicy(self.path_button.sizePolicy().Policy.Expanding, self.path_button.sizePolicy().Policy.Fixed)
        paths.addWidget(self.path_button, 1)
        self.refresh_button = _button("↻", self.refresh_folder)
        self.refresh_button.setToolTip("Обновить список")
        paths.addWidget(self.refresh_button)
        self.hidden = QCheckBox("Скрытые")
        self.hidden.toggled.connect(self.refresh_folder)
        paths.addWidget(self.hidden)
        col.addLayout(paths)
        split = QSplitter(Qt.Horizontal)
        files_panel = QWidget()
        files_col = QVBoxLayout(files_panel)
        files_col.setContentsMargins(0, 0, 0, 0)
        files_col.setSpacing(9)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Найти в этой папке…")
        self.search.setClearButtonEnabled(True)
        self.search.setAccessibleName("Поиск файлов текущей папки")
        self.search.textChanged.connect(self._filter_files)
        files_col.addWidget(self.search)
        self.files = QListWidget()
        self.files.setAccessibleName("Файлы выбранной папки")
        self.files.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.files.setIconSize(QSize(30, 30))
        self.files.currentItemChanged.connect(self._selection_changed)
        self.files.itemDoubleClicked.connect(self._activate_item)
        files_col.addWidget(self.files, 1)
        self.count_label = _label("Папка не выбрана", "muted")
        files_col.addWidget(self.count_label)
        files_panel.setMinimumWidth(210)
        split.addWidget(files_panel)
        preview = QFrame()
        preview.setObjectName("card")
        preview_col = QVBoxLayout(preview)
        preview_col.setContentsMargins(18, 17, 18, 16)
        preview_col.setSpacing(11)
        self.preview_title = _label("Начните с документа", "sectionTitle", True)
        preview_col.addWidget(self.preview_title)
        self.preview_info = _label("Файлы остаются в своих папках. Просмотр ничего не изменяет.", "muted", True)
        preview_col.addWidget(self.preview_info)
        self.preview_stack = QStackedWidget()
        self.empty_page = QWidget()
        empty = QVBoxLayout(self.empty_page)
        empty.addStretch()
        symbol = _label("▤", "emptySymbol")
        symbol.setAlignment(Qt.AlignCenter)
        empty.addWidget(symbol)
        self.empty_message = _label("Откройте папку или перетащите файл в это окно.", "muted", True)
        self.empty_message.setAlignment(Qt.AlignCenter)
        empty.addWidget(self.empty_message)
        shortcuts = QHBoxLayout()
        shortcuts.addStretch()
        shortcuts.addWidget(_button("Загрузки", lambda: self.set_folder(Path.home() / "Downloads")))
        shortcuts.addWidget(_button("Документы", lambda: self.set_folder(Path.home() / "Documents")))
        shortcuts.addStretch()
        empty.addLayout(shortcuts)
        self.unlock_button = _button("Открыть с паролем…", self.unlock_pdf)
        self.unlock_button.hide()
        empty.addWidget(self.unlock_button, 0, Qt.AlignCenter)
        empty.addStretch()
        self.preview_stack.addWidget(self.empty_page)
        self.picture_scroll = QScrollArea()
        self.picture_scroll.setObjectName("documentCanvas")
        self.picture_scroll.setAlignment(Qt.AlignCenter)
        self.picture = QLabel()
        self.picture.setAlignment(Qt.AlignCenter)
        self.picture_scroll.setWidget(self.picture)
        self.preview_stack.addWidget(self.picture_scroll)
        self.text = QPlainTextEdit()
        self.text.setReadOnly(True)
        self.text.setAccessibleName("Текст документа")
        self.text.document().modificationChanged.connect(self._text_modified)
        self.preview_stack.addWidget(self.text)
        preview_col.addWidget(self.preview_stack, 1)
        self.pdf_controls = QWidget()
        self.pdf_controls.hide()
        page_row = QHBoxLayout(self.pdf_controls)
        page_row.setContentsMargins(0, 0, 0, 0)
        self.previous_button = _button("←", lambda: self.page.setValue(self.page.value() - 1))
        self.next_button = _button("→", lambda: self.page.setValue(self.page.value() + 1))
        page_row.addWidget(self.previous_button)
        page_row.addWidget(_label("Страница", "muted"))
        self.page = QSpinBox()
        self.page.setRange(1, 1)
        self.page.valueChanged.connect(self._render_pdf)
        self.page.setAccessibleName("Страница предпросмотра PDF")
        page_row.addWidget(self.page)
        self.page_total = _label("из 1", "muted")
        page_row.addWidget(self.page_total)
        page_row.addWidget(self.next_button)
        page_row.addStretch()
        self.fit_button = _button("По ширине", self.render_preview)
        page_row.addWidget(self.fit_button)
        preview_col.addWidget(self.pdf_controls)
        action_row = QHBoxLayout()
        self.edit_button = _button("Открыть редактор", self.edit_selected, "primary")
        action_row.addWidget(self.edit_button)
        self.save_text_button = _button("Сохранить копию…", self.save_text_copy)
        self.save_text_button.hide()
        action_row.addWidget(self.save_text_button)
        self.tools_button = _button("PDF-инструменты", self.open_tools)
        action_row.addWidget(self.tools_button)
        self.convert_button = _button("Преобразовать", self.open_convert)
        action_row.addWidget(self.convert_button)
        preview_col.addLayout(action_row)
        minor_row = QHBoxLayout()
        self.reveal_button = _button("Показать в Finder", self.reveal_current, "quiet")
        minor_row.addWidget(self.reveal_button)
        minor_row.addStretch()
        self.external_button = _button("Открыть в другой программе ↗", self.open_external, "quiet")
        minor_row.addWidget(self.external_button)
        preview_col.addLayout(minor_row)
        preview.setMinimumWidth(340)
        split.addWidget(preview)
        split.setSizes([260, 680])
        split.setStretchFactor(1, 1)
        col.addWidget(split, 1)
        self.message = _label("", "warning", True)
        self.message.hide()
        col.addWidget(self.message)
        self.preview_timer = QTimer(self)
        self.preview_timer.setSingleShot(True)
        self.preview_timer.setInterval(180)
        self.preview_timer.timeout.connect(self.render_preview)
        self._refresh_buttons()
        if self.store is not None:
            saved = self.store.value("library/last_folder", "")
            if saved and Path(str(saved)).is_dir():
                self.set_folder(saved, remember=False)

    def show_message(self, text):
        self.message.setText(str(text))
        self.message.setVisible(bool(text))

    def choose_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "Открыть папку документов", str(self.folder or Path.home() / "Documents"))
        if folder:
            self.set_folder(folder)

    def choose_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "Открыть документ", str(self.folder or Path.home() / "Downloads"), "Документы и изображения (*.pdf *.docx *.xlsx *.pptx *.odt *.txt *.md *.csv *.json *.xml *.html *.htm *.png *.jpg *.jpeg *.webp *.heic *.heif *.tif *.tiff *.avif *.bmp *.gif *.svg *.epub);;Все файлы (*)")
        if paths:
            self.open_paths(paths)

    def open_paths(self, paths):
        if not paths:
            return False
        first = Path(paths[0]).expanduser()
        if first.is_dir():
            return self.set_folder(first)
        if not first.is_file():
            self.show_message("Этот файл перемещён или больше недоступен.")
            return False
        if not self.set_folder(first.parent, select=first):
            return False
        requested = {str(Path(path).resolve()) for path in paths}
        for index in range(self.files.count()):
            item = self.files.item(index)
            item.setSelected(item.data(Qt.UserRole) in requested)
        if self.current_path != first.resolve():
            self.preview_file(first)
        return True

    def set_folder(self, folder, select=None, remember=True):
        folder = Path(folder).expanduser().resolve()
        if not self.confirm_leave():
            return False
        if not folder.is_dir():
            self.show_message("Эта папка недоступна. Выберите другую.")
            return False
        try:
            entries, truncated = [], False
            with os.scandir(folder) as iterator:
                for entry in iterator:
                    if len(entries) >= MAX_FOLDER_ENTRIES:
                        truncated = True
                        break
                    if entry.name.startswith(".") and not self.hidden.isChecked():
                        continue
                    try:
                        is_dir = entry.is_dir(follow_symlinks=True)
                        if is_dir or entry.is_file(follow_symlinks=True):
                            entries.append((entry.name, str(folder / entry.name), is_dir))
                    except OSError:
                        continue
            entries.sort(key=lambda item: (not item[2], item[0].casefold()))
        except OSError as exc:
            self.show_message(f"Не удалось прочитать папку: {exc}")
            return False
        self._close_pdf()
        self.folder = folder
        self._text_editing = False
        self.text.setReadOnly(True)
        self._entries = entries
        self._selecting = True
        self.files.clear()
        self.current_path = None
        for name, path, directory in entries:
            item = QListWidgetItem(self._icons.icon(QFileInfo(path)), name)
            item.setData(Qt.UserRole, path)
            item.setData(Qt.UserRole + 1, directory)
            item.setToolTip(path)
            item.setSizeHint(QSize(180, 48))
            self.files.addItem(item)
        self._selecting = False
        self.path_button.setText(str(folder))
        self.path_button.setToolTip(str(folder))
        self.search.clear()
        self._filter_files()
        self.show_message(f"Показаны первые {MAX_FOLDER_ENTRIES} объектов. Откройте более узкую папку." if truncated else "")
        self._show_empty("Выберите документ в списке слева.")
        if remember and self.store is not None:
            self.store.setValue("library/last_folder", str(folder))
        if select:
            self._select_path(select)
        self._refresh_buttons()
        return True

    def refresh_folder(self, *_):
        if self.folder:
            self.set_folder(self.folder, select=self.current_path, remember=False)

    def parent_folder(self):
        if self.folder:
            self.set_folder(self.folder.parent)

    def _select_path(self, path):
        path = str(Path(path).resolve())
        for index in range(self.files.count()):
            if self.files.item(index).data(Qt.UserRole) == path:
                self.files.setCurrentRow(index)
                return

    def _filter_files(self, *_):
        query = self.search.text().casefold().strip()
        visible = 0
        for index in range(self.files.count()):
            item = self.files.item(index)
            show = query in item.text().casefold()
            item.setHidden(not show)
            visible += show
        self.count_label.setText(f"Объектов: {visible}" + (f" из {self.files.count()}" if query else ""))

    def selected_files(self):
        selected = [item.data(Qt.UserRole) for item in self.files.selectedItems() if not item.data(Qt.UserRole + 1) and not item.isHidden()]
        return selected or ([str(self.current_path)] if self.current_path and self.current_path.is_file() else [])

    def _selection_changed(self, current, previous):
        if self._selecting or current is None:
            return
        path = Path(current.data(Qt.UserRole))
        if not self.confirm_leave():
            self._selecting = True
            self.files.setCurrentItem(previous)
            self._selecting = False
            return
        if current.data(Qt.UserRole + 1):
            self._text_editing = False
            self.text.setReadOnly(True)
            self.current_path = path
            self._close_pdf()
            self._show_empty("Дважды нажмите на папку, чтобы открыть её.")
            self.preview_title.setText(path.name)
            self.preview_info.setText("Папка")
            self._refresh_buttons()
            return
        self.preview_file(path, confirm=False)

    def _activate_item(self, item):
        path = Path(item.data(Qt.UserRole))
        if item.data(Qt.UserRole + 1):
            self.set_folder(path)
        else:
            self.edit_selected()

    def preview_file(self, path, confirm=True):
        path = Path(path).expanduser().resolve()
        if confirm and not self.confirm_leave():
            return False
        if not path.is_file():
            self.show_message("Файл недоступен: возможно, он перемещён или удалён.")
            return False
        self._close_pdf()
        self.current_path = path
        self._password = ""
        self.page.blockSignals(True)
        self.page.setValue(1)
        self.page.blockSignals(False)
        self._text_editing = False
        self._text_complete = False
        self.text.setReadOnly(True)
        self.text.document().setModified(False)
        self.preview_title.setText(path.name)
        self.preview_info.setText(f"{path.suffix.lstrip('.').upper() or 'ФАЙЛ'}  ·  {human_size(file_size(path))}")
        self.show_message("")
        self.file_selected.emit(str(path))
        if self._processing:
            self._show_empty("Предпросмотр продолжится после обработки файлов.")
        else:
            self.render_preview()
        self._refresh_buttons()
        return True

    def _show_empty(self, message):
        self._preview_kind = "empty"
        self.empty_message.setText(message)
        self.preview_stack.setCurrentWidget(self.empty_page)
        self.pdf_controls.hide()
        self.unlock_button.hide()

    def _close_pdf(self):
        if self._pdf is not None:
            self._pdf.close()
            self._pdf = None

    def render_preview(self):
        if self._processing or not self.current_path or not self.current_path.is_file() or self._text_editing:
            return
        path = self.current_path
        try:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                if self._pdf is None:
                    self._pdf = fitz.open(path)
                    if self._pdf.needs_pass and not self._pdf.authenticate(self._password):
                        self._close_pdf()
                        self._show_empty("PDF защищён паролем. Откройте его для просмотра.")
                        self.unlock_button.show()
                        return
                    if self._pdf.page_count == 0:
                        raise UserError("В этом PDF нет страниц.")
                    previous_page = self.page.value()
                    self.page.blockSignals(True)
                    self.page.setRange(1, self._pdf.page_count)
                    self.page.setValue(min(previous_page, self._pdf.page_count))
                    self.page.blockSignals(False)
                    self.page_total.setText(f"из {self._pdf.page_count}")
                self._preview_kind = "pdf"
                self._render_pdf()
            elif suffix in IMAGE_EXTS:
                from PIL import Image, ImageOps
                with Image.open(path) as source:
                    if source.width * source.height > 50_000_000:
                        raise UserError("Изображение слишком велико для быстрого просмотра. Откройте его в редакторе.")
                    source.draft("RGB", (1600, 1600))
                    image = ImageOps.exif_transpose(source)
                    image.thumbnail((1600, 1600))
                    image = image.convert("RGBA")
                    qimage = QImage(image.tobytes(), image.width, image.height, image.width * 4, QImage.Format_RGBA8888).copy()
                self._preview_kind = "image"
                self._show_picture(QPixmap.fromImage(qimage), fit=True)
            elif suffix in TEXT_EXTS:
                text, self._text_complete, self._text_encoding = read_text_preview(path)
                self._preview_kind = "text"
                self.text.setPlainText(text)
                self.text.document().setModified(False)
                self.preview_stack.setCurrentWidget(self.text)
                self.pdf_controls.hide()
                if not self._text_complete:
                    self.show_message("Показано начало большого файла. Редактирование этого предпросмотра отключено.")
            elif suffix in OFFICE_EXTS:
                self._preview_kind = "office"
                self.text.setPlainText(office_text_preview(path))
                self.text.document().setModified(False)
                self.preview_stack.setCurrentWidget(self.text)
                self.pdf_controls.hide()
                self.show_message("Текстовый предпросмотр. Для PDF-редактора создаётся копия с базовой вёрсткой; исходный документ сохраняется.")
            else:
                self._show_empty("Для этого формата нет встроенного предпросмотра. Можно открыть файл в другой программе.")
            self._refresh_buttons()
        except Exception as exc:
            self._close_pdf()
            self._show_empty("Не удалось показать документ.")
            self.show_message(str(exc))
            self._refresh_buttons()

    def _render_pdf(self, *_):
        if self._processing or self._pdf is None:
            return
        page = self._pdf[self.page.value() - 1]
        width, height = page.rect.width, page.rect.height
        if not all(math.isfinite(value) and value > 0 for value in (width, height)):
            raise UserError("Некорректный размер страницы.")
        desired_width = max(300, self.picture_scroll.viewport().width() - 24)
        scale = min(desired_width / width, 2.0, 1800 / width, 2200 / height, math.sqrt(3_000_000 / (width * height)))
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        image = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format_RGB888).copy()
        self._show_picture(QPixmap.fromImage(image), fit=False)
        self.pdf_controls.show()
        self.previous_button.setEnabled(self.page.value() > 1)
        self.next_button.setEnabled(self.page.value() < self.page.maximum())

    def _show_picture(self, pixmap, fit=False):
        if fit:
            size = self.picture_scroll.viewport().size() - QSize(24, 24)
            if size.width() > 0 and size.height() > 0:
                pixmap = pixmap.scaled(size, Qt.KeepAspectRatio, Qt.SmoothTransformation)
        self.picture.setPixmap(pixmap)
        self.picture.setFixedSize(pixmap.size())
        self.preview_stack.setCurrentWidget(self.picture_scroll)
        self.pdf_controls.setVisible(self._preview_kind == "pdf")

    def unlock_pdf(self):
        if self._processing:
            return
        password, ok = QInputDialog.getText(self, "Открыть PDF", "Пароль документа", QLineEdit.Password)
        if ok:
            self._password = password
            self.render_preview()
            if self._pdf is None:
                self.show_message("Пароль не подошёл. Повторите ввод.")

    def _refresh_buttons(self):
        file = self.current_path if self.current_path and self.current_path.is_file() else None
        suffix = file.suffix.lower() if file else ""
        allowed = bool(file) and not self._processing
        editable = suffix in IMAGE_EXTS | OFFICE_EXTS | {".pdf"} or (suffix in TEXT_EXTS and self._text_complete)
        self.edit_button.setEnabled(allowed and editable and not self._text_editing)
        self.edit_button.setText("Редактировать текст" if suffix in TEXT_EXTS else "Создать PDF и редактировать" if suffix in OFFICE_EXTS else "Открыть редактор")
        self.save_text_button.setVisible(self._text_editing)
        self.save_text_button.setEnabled(allowed and self.text.document().isModified())
        self.tools_button.setEnabled(allowed)
        self.convert_button.setEnabled(allowed)
        self.reveal_button.setEnabled(bool(self.current_path))
        self.external_button.setEnabled(bool(file))
        self.up_button.setEnabled(bool(self.folder) and self.folder != self.folder.parent)
        self.refresh_button.setEnabled(bool(self.folder) and not self._processing)
        self.open_files_button.setEnabled(not self._processing)
        self.choose_folder_button.setEnabled(not self._processing)
        self.path_button.setEnabled(not self._processing)
        self.hidden.setEnabled(not self._processing)

    def edit_selected(self):
        if self._processing or not self.current_path or not self.current_path.is_file():
            return
        path = self.current_path
        try:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                from .pdf_workspace import PdfEditor
                editor = PdfEditor(str(path), self.controller)
                from .theme import register_window
                register_window(editor)
                self.controller.open_window(editor)
            elif suffix in IMAGE_EXTS:
                from .radial import Item
                self.controller.run_item(Item("Edit", key="tool:Edit"), [str(path)])
            elif suffix in TEXT_EXTS and self._text_complete:
                self._text_editing = True
                self.text.setReadOnly(False)
                self.text.setFocus()
                self.show_message("Изменения будут сохранены отдельной копией. Исходный файл останется прежним.")
            elif suffix in OFFICE_EXTS:
                from .document_services import office_to_pdf
                self._pending_edit_source = str(path)
                self.controller.submit(f"library:{PREPARE_EDIT_ACTION}", [str(path)], lambda: office_to_pdf(path), working="Создаём PDF для редактора…")
            else:
                self.show_message("Для этого формата используйте преобразование или внешнюю программу.")
            self._refresh_buttons()
        except Exception as exc:
            self.show_message(str(exc))

    def operation_finished(self, result):
        if self._pending_edit_source and result.title == PREPARE_EDIT_ACTION and self._pending_edit_source in {str(path) for path in result.srcs}:
            outputs = [str(path) for path in result.outs if Path(path).suffix.lower() == ".pdf"]
            self._pending_edit_source = None
            if outputs:
                self._pending_editor = outputs[0]
                if not self._processing:
                    self._open_pending_editor()
        elif self.folder and any(Path(path).parent == self.folder for path in result.outs):
            # Do not replace an active editor's unsaved text merely to refresh the list.
            if not self._text_editing:
                self.refresh_folder()

    def _open_pending_editor(self):
        path = self._pending_editor
        self._pending_editor = None
        if path and self.open_paths([path]):
            self.edit_selected()

    def open_tools(self):
        if not self._processing:
            self.tools_requested.emit(self.selected_files())

    def open_convert(self):
        if not self._processing:
            self.convert_requested.emit(self.selected_files())

    def _text_modified(self, *_):
        if hasattr(self, "save_text_button"):
            self._refresh_buttons()

    def save_text_copy(self):
        if self._processing or not self._text_editing or not self.current_path:
            return False
        destination, _ = QFileDialog.getSaveFileName(self, "Сохранить копию документа", str(sibling(self.current_path, "Edited")), "Все файлы (*)")
        if not destination:
            return False
        try:
            output = self.write_text_copy(destination)
            self.controller.notify("library:Редактирование текста", [self.current_path], [output])
            self.text.document().setModified(False)
            self.show_message(f"Копия сохранена: {output.name}")
            return True
        except Exception as exc:
            self.show_message(str(exc))
            return False

    def write_text_copy(self, destination):
        if not self._text_editing or not self._text_complete or not self.current_path:
            raise UserError("Сначала откройте небольшой текстовый документ для редактирования.")
        output = Path(destination).expanduser()
        if output.resolve() == self.current_path.resolve():
            raise UserError("Выберите другое имя: исходный документ сохраняется без изменений.")
        # UTF-8 keeps every character typed in the editor, including Cyrillic.
        data = self.text.toPlainText().encode("utf-8")
        atomic_save(lambda path: Path(path).write_bytes(data), output)
        return output

    def confirm_leave(self):
        if self._text_editing and self.text.document().isModified():
            answer = QMessageBox.question(self, "Сохранить изменения?", "Сохранить изменённый текст отдельной копией?", QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel, QMessageBox.Save)
            if answer == QMessageBox.Cancel:
                return False
            if answer == QMessageBox.Save and not self.save_text_copy():
                return False
            self.text.document().setModified(False)
        return True

    def set_processing(self, busy):
        busy = bool(busy)
        if busy == self._processing:
            return
        self._processing = busy
        if busy:
            self.preview_timer.stop()
            self._close_pdf()
            if self._preview_kind == "pdf":
                self.preview_info.setText("Предпросмотр приостановлен на время обработки файлов.")
        else:
            if self.current_path and self.current_path.is_file():
                self.preview_info.setText(f"{self.current_path.suffix.lstrip('.').upper() or 'ФАЙЛ'}  ·  {human_size(file_size(self.current_path))}")
            if self._pending_editor:
                self._open_pending_editor()
            elif self.current_path and not self._text_editing:
                self.render_preview()
        self.files.setEnabled(not busy)
        self.text.setReadOnly(busy or not self._text_editing)
        self.pdf_controls.setEnabled(not busy)
        self._refresh_buttons()

    def reveal_current(self):
        if not self.current_path:
            return
        if sys.platform == "darwin":
            subprocess.Popen(["open", "-R", str(self.current_path)])
        else:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.current_path.parent)))

    def open_external(self):
        if self.current_path and not QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.current_path))):
            self.show_message("Не удалось открыть документ. Используйте Finder.")

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "preview_timer") and not self._processing and self._preview_kind in ("pdf", "image"):
            self.preview_timer.start()

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls() and any(url.isLocalFile() for url in event.mimeData().urls()):
            event.acceptProposedAction()

    def dropEvent(self, event):
        paths = [url.toLocalFile() for url in event.mimeData().urls() if url.isLocalFile()]
        if paths and not self._processing:
            self.open_paths(paths)
            event.acceptProposedAction()
