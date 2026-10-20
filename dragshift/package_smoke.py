"""Executable-bundle smoke checks with ephemeral documents and no user-state writes.

Called by the app's explicit self-test mode. No clipboard, preferences, watcher,
network, scanner, screen capture, printing, Keychain or installed app is modified.
"""
from __future__ import annotations

import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path


def run() -> dict:
    from PIL import Image, ImageDraw, ImageFont
    import pikepdf
    from . import convert, document_services as native, pdf_ops, pdf_signing, pdf_standards
    from .util import CONFIG

    checks = []

    def record(name, condition=True):
        if not condition:
            raise RuntimeError(f"Package smoke failed: {name}")
        checks.append(name)

    def readable_pdf(path, expected_text=None):
        with convert.fitz.open(path) as doc:
            if doc.needs_pass or not doc.page_count:
                return False
            doc[0].get_pixmap(matrix=convert.fitz.Matrix(.3, .3))
            return expected_text is None or expected_text in "".join(p.get_text() for p in doc)

    # Runtime CONFIG is an in-memory dictionary, never the user's QSettings file.
    previous_output = CONFIG.get("output", "next")
    CONFIG["output"] = "next"
    try:
        with tempfile.TemporaryDirectory(prefix="dna-package-smoke-") as directory:
            root = Path(directory)
            source = root / "Проверка PDF.pdf"
            with convert.fitz.open() as doc:
                page = doc.new_page(width=400, height=300)
                page.insert_text((25, 28), "DNA package smoke", fontsize=14)
                # A real ruled table exercises the PDF table detector/XLSX exporter.
                for x in (30, 180, 330):
                    page.draw_line((x, 60), (x, 180), color=(0, 0, 0))
                for y in (60, 100, 140, 180):
                    page.draw_line((30, y), (330, y), color=(0, 0, 0))
                for row, entries in enumerate((("Name", "Value"), ("Alpha", "100"), ("Beta", "200"))):
                    for col, value in enumerate(entries):
                        page.insert_text((40 + col * 150, 84 + row * 40), value, fontsize=12)
                doc.new_page(width=400, height=300).insert_text((25, 28), "Second page")
                doc.set_metadata({"title": "PACKAGE-ORIGINAL"})
                doc.save(source)
            original = source.read_bytes()
            rotated = pdf_ops.rotate_pages(source, "2")[0]
            record("PDF operations", readable_pdf(rotated, "Second page"))
            reordered = pdf_ops.reorder_pages(source, "2,1,2")[0]
            with convert.fitz.open(reordered) as doc:
                record("Independent duplicated pages", doc.page_count == 3 and doc[0].xref != doc[2].xref)
            marked = pdf_ops.watermark_pdf(source, "ПРОВЕРКА DNA.apps", pages="1")[0]
            record("Embedded Unicode font", readable_pdf(marked, "ПРОВЕРКА"))
            linearized = pdf_ops.optimize_pdf(source, linearize=True)[0]
            with pikepdf.open(linearized) as pdf:
                record("QPDF real linearization", pdf.is_linearized and pdf.check_linearization())

            image = root / "Фото.png"
            Image.new("RGBA", (160, 100), (105, 25, 55, 150)).save(image)
            for target in ("JPG", "WEBP", "TIFF", "PDF"):
                output = convert.image_to(image, target)[0]
                if target == "PDF":
                    record(f"Image → {target}", readable_pdf(output))
                else:
                    with Image.open(output) as result:
                        result.load()
                        record(f"Image → {target}", result.width > 0)

            from cryptography import x509
            from cryptography.hazmat.primitives import hashes, serialization
            from cryptography.hazmat.primitives.asymmetric import rsa
            from cryptography.hazmat.primitives.serialization import pkcs12
            from cryptography.x509.oid import NameOID
            key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
            name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "DNA package smoke")])
            now = datetime.now(timezone.utc)
            certificate = (x509.CertificateBuilder().subject_name(name).issuer_name(name)
                .public_key(key.public_key()).serial_number(x509.random_serial_number())
                .not_valid_before(now - timedelta(days=1)).not_valid_after(now + timedelta(days=1))
                .add_extension(x509.BasicConstraints(ca=True, path_length=None), critical=True)
                .add_extension(x509.KeyUsage(digital_signature=True, content_commitment=True,
                    key_encipherment=False, data_encipherment=False, key_agreement=False,
                    key_cert_sign=True, crl_sign=True, encipher_only=False, decipher_only=False), critical=True)
                .sign(key, hashes.SHA256()))
            p12 = root / "Ephemeral.p12"
            p12.write_bytes(pkcs12.serialize_key_and_certificates(b"smoke", key, certificate, None,
                            serialization.BestAvailableEncryption(b"temporary-test-password")))
            signed = pdf_signing.sign_pdf(source, p12, "temporary-test-password")[0]
            signatures = pdf_signing.inspect_signatures(signed)
            record("PKCS#12 signature integrity", len(signatures) == 1 and signatures[0]["intact"] and signatures[0]["valid"])
            record("Offline trust is distinguished", not signatures[0]["trusted"])
            record("Unicode signature report", readable_pdf(pdf_signing.signature_report(signed)[0], "целостность"))
            # Deliberately tamper signed bytes without changing file offsets.
            tampered = root / "Tampered.pdf"
            tampered.write_bytes(signed.read_bytes().replace(b"PACKAGE-ORIGINAL", b"PACKAGE-MODIFIED", 1))
            invalid = pdf_signing.inspect_signatures(tampered)
            record("Signature tampering detected", len(invalid) == 1 and not invalid[0]["intact"])

            for kind, fn in (("PDF/A", pdf_standards.convert_pdfa), ("PDF/X", pdf_standards.convert_pdfx)):
                result = fn(source)[0]
                evidence = pdf_standards.inspect_standard_pdf(result)
                record(f"Ghostscript {kind} conversion", evidence["all_fonts_embedded"] and bool(evidence["output_intents"]) and readable_pdf(result))

            if sys.platform in {"darwin", "win32"} and getattr(sys, "frozen", False):
                import subprocess
                from .office_renderer import find_libreoffice, _renderer_environment
                from .runtime_process import external_run
                engine = find_libreoffice()
                bundled_runtime = (Path(sys._MEIPASS) / "libreoffice" if sys.platform == "win32" else
                                   Path(sys.executable).resolve().parent.parent / "Resources" / "libreoffice")
                record("LibreOffice resolves inside app bundle", engine is not None
                       and engine.is_relative_to(bundled_runtime.resolve()))
                version = external_run(
                    [str(engine), "-env:UserInstallation=" + (root / "version-profile").as_uri(),
                     "--headless", "--version"],
                    cwd=root, env=_renderer_environment(), stdin=subprocess.DEVNULL,
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=30,
                )
                record("Bundled LibreOffice executable launches", version.returncode == 0
                       and "LibreOffice" in version.stdout)

            from openpyxl import Workbook, load_workbook
            workbook = Workbook()
            workbook.active.append(["Name", "Value"])
            workbook.active.append(["Alpha", 123])
            xlsx = root / "Sheet.xlsx"
            workbook.save(xlsx)
            workbook.close()
            record("XLSX → PDF", readable_pdf(native.office_to_pdf(xlsx)[0], "Alpha"))
            exported_xlsx = native.export_document(source, "xlsx")[0]
            workbook = load_workbook(exported_xlsx, read_only=True)
            try:
                record("PDF table → XLSX", any("Alpha" in str(cell.value) for row in workbook.active for cell in row))
            finally:
                workbook.close()

            from pptx import Presentation
            from pptx.util import Inches
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Smoke presentation"
            pptx = root / "Slides.pptx"
            presentation.save(pptx)
            record("PPTX → PDF", readable_pdf(native.office_to_pdf(pptx)[0], "Smoke presentation"))
            exported_pptx = native.export_document(source, "pptx", dpi=72)[0]
            record("PDF → PPTX", len(Presentation(exported_pptx).slides) == 2)

            from docx import Document
            from docx.enum.section import WD_ORIENT, WD_SECTION_START
            from docx.shared import Mm, Pt, RGBColor
            document = Document()
            document.styles["Normal"].font.name = "Arial"
            document.styles["Normal"].font.size = Pt(11)
            portrait = document.sections[0]
            portrait.page_width, portrait.page_height = Mm(210), Mm(297)
            portrait.left_margin = portrait.right_margin = Mm(20)
            portrait.top_margin = portrait.bottom_margin = Mm(20)
            portrait.header.paragraphs[0].text = "Portrait header"
            portrait.footer.paragraphs[0].text = "Portrait footer"
            heading = document.add_paragraph().add_run("Документ DNA smoke")
            heading.bold = True
            heading.font.size = Pt(18)
            heading.font.color.rgb = RGBColor(0x9C, 0x24, 0x3E)
            table = document.add_table(rows=2, cols=3)
            table.style = "Table Grid"
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged table heading"
            table.cell(0, 2).text = "Total"
            for cell, text in zip(table.rows[1].cells, ("Item", "Count", "Value")):
                cell.text = text
            document.add_page_break()
            document.add_paragraph("Explicit page break marker")
            landscape = document.add_section(WD_SECTION_START.NEW_PAGE)
            landscape.orientation = WD_ORIENT.LANDSCAPE
            landscape.page_width, landscape.page_height = Mm(297), Mm(210)
            landscape.left_margin = landscape.right_margin = Mm(15)
            landscape.top_margin = landscape.bottom_margin = Mm(15)
            landscape.header.is_linked_to_previous = False
            landscape.footer.is_linked_to_previous = False
            landscape.header.paragraphs[0].text = "Landscape header"
            landscape.footer.paragraphs[0].text = "Landscape footer"
            document.add_paragraph("Landscape section marker")
            docx = root / "Document.docx"
            document.save(docx)
            original_docx = docx.read_bytes()
            exported_docx = native.office_to_pdf(docx)[0]
            record("DOCX → PDF", readable_pdf(exported_docx, "Документ"))
            with convert.fitz.open(exported_docx) as rendered:
                expected_sizes = ((210, 297), (210, 297), (297, 210))
                record("DOCX page sizes and section orientation", rendered.page_count == 3
                       and all(abs(page.rect.width - width * 72 / 25.4) < 1
                               and abs(page.rect.height - height * 72 / 25.4) < 1
                               for page, (width, height) in zip(rendered, expected_sizes)))
                texts = [page.get_text() for page in rendered]
                record("DOCX explicit page and section breaks", "Explicit page break marker" in texts[1]
                       and "Explicit page break marker" not in texts[0]
                       and "Landscape section marker" in texts[2])
                record("DOCX section headers and footers", all("Portrait header" in text and "Portrait footer" in text
                       and "Landscape header" not in text for text in texts[:2])
                       and "Landscape header" in texts[2] and "Landscape footer" in texts[2]
                       and "Portrait header" not in texts[2])
                spans = [span for block in rendered[0].get_text("dict")["blocks"] if "lines" in block
                         for line in block["lines"] for span in line["spans"]]
                record("DOCX bold text size and color", any("DNA smoke" in span["text"]
                       and span["flags"] & 16 and abs(span["size"] - 18) < .5
                       and span["color"] == 0x9C243E for span in spans))
                cells = [rendered[0].search_for(text) for text in ("Merged table heading", "Item", "Count", "Value", "Total")]
                record("DOCX merged table content and columns", all(cells)
                       and cells[1][0].x0 < cells[2][0].x0 < cells[3][0].x0
                       and abs(cells[0][0].x0 - cells[1][0].x0) < 2
                       and abs(cells[4][0].x0 - cells[3][0].x0) < 2)
            record("DOCX source unchanged", docx.read_bytes() == original_docx)
            record("PDF → DOCX", bool(Document(convert.pdf_to_docx(source)[0]).paragraphs))
            markup = root / "Local.html"
            markup.write_text("<h1>Проверка HTML</h1><p>Local static content</p>", encoding="utf-8")
            record("Static HTML → PDF", readable_pdf(native.html_to_pdf(markup)[0], "Проверка HTML"))

            if sys.platform == "darwin":
                ocr_image = Image.new("RGB", (1400, 400), "white")
                font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 56)
                drawing = ImageDraw.Draw(ocr_image)
                drawing.text((60, 65), "DNA.apps 2026", font=font, fill="black")
                drawing.text((60, 180), "Проверка документа", font=font, fill="black")
                ocr_path = root / "Скан.png"
                ocr_image.save(ocr_path)
                recognized = "\n".join(line.text for line in native.recognize_image(ocr_image))
                record("Real Apple Vision Russian/English OCR", "DNA.apps" in recognized and "Проверка" in recognized)
                searchable = native.ocr_document(ocr_path)[0]
                record("Searchable OCR PDF", readable_pdf(searchable, "Проверка"))
            elif sys.platform == "win32":
                import winrt.windows.foundation
                import winrt.windows.foundation.collections
                import winrt.windows.globalization
                import winrt.windows.graphics.imaging
                import winrt.windows.media.ocr
                import winrt.windows.storage
                import winrt.windows.storage.streams
                record("Windows OCR projections packaged (recognition needs Windows language packs)")

            from PySide6.QtCore import QPointF, QCoreApplication, QEvent
            from PySide6.QtWidgets import QApplication
            from .desktop import DesktopWindow
            from .pdf_workspace import PdfEditor, TOOLS
            from .tray_panel import TrayPanel
            app = QApplication.instance() or QApplication([])

            class EphemeralController:
                settings = {"output": "next"}
                active_jobs = 0
                def trigger_label(self):
                    return "Shift"
                def build_rings(self, paths):
                    return {}
                def __getattr__(self, name):
                    return lambda *args, **kwargs: None

            controller = EphemeralController()
            desktop = DesktopWindow(controller)
            tray = TrayPanel(controller)
            editor = PdfEditor(str(source), controller)
            try:
                record("Desktop/PDF catalog construction", len(TOOLS) == 55 and desktop.pdf_workspace is not None)
                record("Editor-first workspace", desktop.pages.currentIndex() == 3)
                from .theme import apply_theme, tokens
                apply_theme("dark", app)
                record("Dark theme", app.property("docuFreedomTheme") == "dark")
                apply_theme("light", app)
                workspace = desktop.pdf_workspace
                workspace.toggle_favorite("merge")
                record("Pinned tools persistence", "merge" in controller.settings.get("favorite_tools", []))
                workspace.toggle_favorite("merge")
                # Hidden workspaces deliberately release their raster previews.
                # Exercise the visible UI path before asserting decoded pixels.
                desktop.show()
                app.processEvents()
                desktop.preview_file(source)
                app.processEvents()
                # macOS exposes /var as an alias of /private/var. The browser
                # intentionally stores canonical paths, so compare resolved names.
                record("Internal file preview", desktop.pages.currentIndex() == 3
                       and desktop.library.current_path == source.resolve()
                       and desktop.library._preview_kind == "pdf"
                       and not desktop.library.picture.pixmap().isNull())
                tray.result.set_path(marked)
                record("Tray latest-result card", tray.result.path == marked)
                editor.insert_text(QPointF(25, 215), "Добавленный текст", fontsize=12)
                saved = editor.write_copy(root / "Editor copy.pdf")
                record("Visual editor Unicode save", readable_pdf(saved, "Добавленный текст"))
                editor.undo()
                record("Visual editor undo", "Добавленный текст" not in editor.doc[0].get_text())
            finally:
                editor.dirty = False
                editor.close()
                tray.close()
                desktop.library.set_processing(True)  # close the ephemeral PDF before temp cleanup
                desktop.close()
                editor.deleteLater()
                tray.deleteLater()
                desktop.deleteLater()
                app.processEvents()
                QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
            record("Original files preserved", source.read_bytes() == original)
            record("Atomic saves cleaned up", not list(root.glob(".~*")))
    finally:
        CONFIG["output"] = previous_output
    return {"status": "ok", "count": len(checks), "checks": checks}
