"""Local document services used by the desktop PDF workspace.

OCR uses Apple's on-device Vision framework. Office and HTML conversions are
intentionally basic layout exporters, not a substitute for an Office/browser
printing engine. No file is uploaded by any service in this module.
"""
from __future__ import annotations

import datetime as dt
import html
import io
import math
import ssl
import subprocess
import sys
import tempfile
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pymupdf as fitz
from PIL import Image, ImageOps, ImageSequence

from .convert import atomic_save, docx_to_pdf, flatten, open_pdf
from .util import IMAGE_EXTS, UserError, sibling, unique_path

OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".odt"}
OFFICE_LIMITATIONS = (
    "Базовая вёрстка: текст, таблицы и изображения. Диаграммы, SmartArt, макросы "
    "и сложные эффекты не поддерживаются. Excel использует сохранённые значения "
    "формул; формулы без сохранённого результата выводятся как текст."
)
HTML_LIMITATIONS = "Статическая HTML-страница без JavaScript; сложная вёрстка сайта может отличаться."
MAX_HTML_BYTES = 20 * 1024 * 1024
MAX_ZIP_UNCOMPRESSED = 256 * 1024 * 1024
MAX_OCR_PIXELS = 24_000_000


@dataclass(frozen=True)
class OCRLine:
    text: str
    # Normalized origin at the image's top-left, unlike Vision's bottom-left.
    x: float
    y: float
    width: float
    height: float
    confidence: float


def recognize_image(image: Image.Image, languages=("ru-RU", "en-US")) -> list[OCRLine]:
    """Recognize one image locally, retaining each line's placement."""
    if sys.platform != "darwin":
        raise UserError("Распознавание Apple Vision доступно только в macOS.")
    try:
        import Vision
        import objc
        from Foundation import NSData
    except ImportError as exc:
        raise UserError("В этой сборке недоступен модуль Apple Vision OCR.") from exc
    if image.width * image.height > MAX_OCR_PIXELS:
        image = image.copy()
        scale = math.sqrt(MAX_OCR_PIXELS / (image.width * image.height))
        image.thumbnail((int(image.width * scale), int(image.height * scale)), Image.Resampling.LANCZOS)
    raw = io.BytesIO()
    flatten(image).save(raw, "PNG")
    payload = raw.getvalue()
    with objc.autorelease_pool():
        request = Vision.VNRecognizeTextRequest.alloc().init()
        request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
        supported, error = request.supportedRecognitionLanguagesAndReturnError_(None)
        if error:
            raise UserError(f"Не удалось получить языки OCR: {error.localizedDescription()}")
        selected = [language for language in languages if language in supported]
        if not selected:
            raise UserError("Указанные языки OCR не поддерживаются этой версией macOS.")
        request.setRecognitionLanguages_(selected)
        request.setUsesLanguageCorrection_(True)
        data = NSData.dataWithBytes_length_(payload, len(payload))
        handler = Vision.VNImageRequestHandler.alloc().initWithData_options_(data, {})
        success, error = handler.performRequests_error_([request], None)
        if not success:
            detail = error.localizedDescription() if error else "неизвестная ошибка"
            raise UserError(f"Не удалось распознать изображение: {detail}")
        lines = []
        for observation in request.results() or []:
            candidates = observation.topCandidates_(1)
            if not candidates:
                continue
            candidate = candidates[0]
            rect = observation.boundingBox()
            lines.append(OCRLine(str(candidate.string()), float(rect.origin.x),
                                 float(1 - rect.origin.y - rect.size.height),
                                 float(rect.size.width), float(rect.size.height),
                                 float(candidate.confidence())))
    return lines


def _ocr_pages(src: Path, dpi: int):
    if src.suffix.lower() == ".pdf":
        with open_pdf(src) as document:
            for page in document:
                actual_dpi = min(dpi, 72 * math.sqrt(MAX_OCR_PIXELS / (page.rect.width * page.rect.height)))
                pixmap = page.get_pixmap(matrix=fitz.Matrix(actual_dpi / 72, actual_dpi / 72), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                yield image, page.rect.width, page.rect.height, document.page_count
    elif src.suffix.lower() in IMAGE_EXTS:
        with Image.open(src) as source:
            count = getattr(source, "n_frames", 1)
            for frame in ImageSequence.Iterator(source):
                image = flatten(ImageOps.exif_transpose(frame.copy()).convert("RGBA"))
                yield image, image.width * 72 / dpi, image.height * 72 / dpi, count
    else:
        raise UserError("Для OCR выберите PDF или изображение.")


def ocr_document(src: str | Path, *, output: str = "pdf", languages=("ru-RU", "en-US"),
                 dpi: int = 200, progress: Callable[[int, int], None] | None = None) -> list[Path]:
    """Create a searchable facsimile PDF or UTF-8 text without replacing the source."""
    if output not in {"pdf", "txt"}:
        raise UserError("Результат OCR должен быть PDF или TXT.")
    if not 72 <= dpi <= 600:
        raise UserError("Разрешение OCR должно быть от 72 до 600 dpi.")
    src = Path(src)
    destination = sibling(src, "OCR", "." + output)
    recognized = []
    total_lines = 0
    document = fitz.open()
    font = fitz.Font("cjk") if output == "pdf" else None
    try:
        for index, (image, width, height, total) in enumerate(_ocr_pages(src, dpi)):
            lines = recognize_image(image, languages)
            total_lines += len(lines)
            recognized.append("\n".join(line.text for line in lines))
            if output == "pdf":
                page = document.new_page(width=width, height=height)
                image_bytes = io.BytesIO()
                image.save(image_bytes, "PNG")
                page.insert_image(page.rect, stream=image_bytes.getvalue())
                page.insert_font(fontname="OCRUnicode", fontbuffer=font.buffer)
                for line in lines:
                    if not line.text.strip() or line.width <= 0 or line.height <= 0:
                        continue
                    size = max(1.0, line.height * height * 0.92)
                    baseline = fitz.Point(line.x * width, (line.y + line.height * 0.88) * height)
                    length = font.text_length(line.text, fontsize=size)
                    horizontal_scale = line.width * width / length if length else 1
                    page.insert_text(baseline, line.text, fontsize=size, fontname="OCRUnicode",
                                     render_mode=3, morph=(baseline, fitz.Matrix(horizontal_scale, 1)))
            if progress:
                progress(index + 1, total)
        if not total_lines:
            raise UserError("Текст не найден. Попробуйте более чёткое изображение или другой язык OCR.")
        if output == "pdf":
            document.set_metadata({"title": src.stem, "producer": "DocuFreedom • Apple Vision OCR"})
            document.subset_fonts()
            atomic_save(lambda target: document.save(target, garbage=4, deflate=True), destination)
        else:
            atomic_save(lambda target: Path(target).write_text("\n\f\n".join(recognized), encoding="utf-8"), destination)
    finally:
        document.close()
    return [destination]


def _render_html(markup: str, destination: Path, *, archive=None, landscape=False, css="") -> None:
    """Paginate static HTML with bundled MuPDF, retaining selectable Unicode text."""
    from .convert import CSS
    page_rect = fitz.paper_rect("a4-l" if landscape else "a4")
    content_rect = page_rect + (36, 36, -36, -36)
    story = fitz.Story(html=markup, user_css=CSS + css, archive=archive)
    page_count = 0

    def rectangles(_number, _filled):
        nonlocal page_count
        page_count += 1
        if page_count > 2000:
            raise UserError("Документ слишком велик или содержит элемент, который не помещается на страницу.")
        return page_rect, content_rect, None

    document = story.write_with_links(rectangles)
    try:
        document.set_metadata({"producer": "DocuFreedom • static document export"})
        atomic_save(lambda target: document.save(target, garbage=4, deflate=True), destination)
    finally:
        document.close()


def _sanitize_html(markup: str, source_dir: Path | None = None):
    """Disable scripts/remote assets. Only embed local images within the HTML folder."""
    from lxml import etree
    from lxml import html as lhtml
    parser = lhtml.HTMLParser(encoding="utf-8", no_network=True)
    try:
        root = lhtml.document_fromstring(markup.encode("utf-8"), parser=parser)
    except (etree.ParserError, ValueError) as exc:
        raise UserError("Не удалось прочитать HTML-файл.") from exc
    for element in root.xpath("//script|//iframe|//object|//embed|//link|//base|//form"):
        element.drop_tree()
    archive = fitz.Archive()
    for index, image in enumerate(root.xpath("//img")):
        value = image.get("src", "")
        parsed = urllib.parse.urlsplit(value)
        if source_dir is None or parsed.scheme or parsed.netloc:
            image.drop_tree()
            continue
        path = (source_dir / urllib.parse.unquote(parsed.path)).resolve()
        if not path.is_relative_to(source_dir.resolve()) or not path.is_file() or path.stat().st_size > MAX_HTML_BYTES:
            image.drop_tree()
            continue
        name = f"image-{index}{path.suffix}"
        archive.add((path.read_bytes(), name))
        image.set("src", name)
    return lhtml.tostring(root, encoding="unicode"), archive


def html_to_pdf(src: str | Path) -> list[Path]:
    """Render a local HTML file; local images from its folder are included."""
    src = Path(src)
    if src.stat().st_size > MAX_HTML_BYTES:
        raise UserError("HTML-файл больше 20 МБ.")
    raw = src.read_bytes()
    from lxml import html as lhtml
    # Let the HTML parser honor its declared encoding, then normalize to Unicode.
    try:
        try:
            raw.decode("utf-8-sig")
            parser = lhtml.HTMLParser(encoding="utf-8", no_network=True)
        except UnicodeDecodeError:
            parser = lhtml.HTMLParser(no_network=True)
        markup = lhtml.tostring(lhtml.document_fromstring(raw, parser=parser), encoding="unicode")
    except Exception as exc:
        raise UserError("Не удалось прочитать HTML-файл.") from exc
    markup, archive = _sanitize_html(markup, src.parent)
    destination = sibling(src, ext=".pdf")
    _render_html(markup, destination, archive=archive)
    return [destination]


class _HTTPOnlyRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        if urllib.parse.urlsplit(newurl).scheme.lower() not in {"http", "https"}:
            raise UserError("Перенаправление на этот тип адреса не поддерживается.")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def webpage_to_pdf(url: str, output_dir: str | Path) -> list[Path]:
    """Fetch the explicitly submitted URL only; no subresources or uploads."""
    parts = urllib.parse.urlsplit(url.strip())
    if parts.scheme.lower() not in {"http", "https"} or not parts.netloc or parts.username or parts.password:
        raise UserError("Введите полный адрес http:// или https:// без логина и пароля.")
    request = urllib.request.Request(url.strip(), headers={"User-Agent": "DocuFreedom/3.0 (static PDF export)"})
    try:
        # Python's framework certificate path is absent in the standalone app.
        # Bundle an explicit CA set rather than disabling HTTPS verification.
        import certifi
        context = ssl.create_default_context(cafile=certifi.where())
        opener = urllib.request.build_opener(_HTTPOnlyRedirect(), urllib.request.HTTPSHandler(context=context))
        with opener.open(request, timeout=30) as response:
            content_type = response.headers.get_content_type()
            if content_type not in {"text/html", "application/xhtml+xml"}:
                raise UserError("Адрес должен вести на HTML-страницу.")
            payload = response.read(MAX_HTML_BYTES + 1)
            if len(payload) > MAX_HTML_BYTES:
                raise UserError("Веб-страница больше 20 МБ.")
            markup = payload.decode(response.headers.get_content_charset() or "utf-8", errors="replace")
    except UserError:
        raise
    except Exception as exc:
        raise UserError(f"Не удалось загрузить страницу: {exc}") from exc
    markup, archive = _sanitize_html(markup)
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    name = "".join(character if character.isalnum() or character in "-." else "_" for character in (parts.hostname or "Webpage"))
    destination = unique_path(directory / f"{name}.pdf")
    _render_html(markup, destination, archive=archive)
    return [destination]


def _check_office_archive(src: Path) -> None:
    try:
        with zipfile.ZipFile(src) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_ZIP_UNCOMPRESSED:
                raise UserError("Документ слишком большой: распакованный размер превышает 256 МБ.")
    except zipfile.BadZipFile as exc:
        raise UserError("Повреждённый или неподдерживаемый документ Office.") from exc


def _xlsx_to_pdf(src: Path, destination: Path) -> None:
    import openpyxl
    values = openpyxl.load_workbook(src, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(src, read_only=True, data_only=False)
    try:
        combined = fitz.open()
        try:
            for sheet in values:
                if sheet.sheet_state != "visible":
                    continue
                if sheet.max_row * sheet.max_column > 200_000:
                    raise UserError("Лист Excel слишком большой для базового экспорта (более 200 000 ячеек).")
                rows = []
                formula_rows = formulas[sheet.title].iter_rows()
                for value_row, formula_row in zip(sheet.iter_rows(), formula_rows):
                    cells = []
                    for cell, formula in zip(value_row, formula_row):
                        value = cell.value if cell.value is not None else formula.value
                        if isinstance(value, (dt.datetime, dt.date, dt.time)):
                            value = value.isoformat(sep=" ") if isinstance(value, dt.datetime) else value.isoformat()
                        text = "" if value is None else str(value)
                        if len(text) > 32767:
                            text = text[:32767]
                        style = "font-weight:bold;" if cell.font and cell.font.bold else ""
                        cells.append(f'<td style="{style}">{html.escape(text).replace(chr(10), "<br>")}</td>')
                    rows.append("<tr>" + "".join(cells) + "</tr>")
                markup = f"<h1>{html.escape(sheet.title)}</h1><table>{''.join(rows)}</table>"
                with tempfile.TemporaryDirectory(prefix="dna-xlsx-") as temp:
                    part = Path(temp) / "sheet.pdf"
                    _render_html(markup, part, landscape=sheet.max_column > 6,
                                 css="table {width:100%;} td {font-size:8pt; word-break:break-all;}")
                    with fitz.open(part) as document:
                        combined.insert_pdf(document)
            if not combined.page_count:
                raise UserError("В книге Excel нет видимых листов.")
            atomic_save(lambda target: combined.save(target, garbage=4, deflate=True), destination)
        finally:
            combined.close()
    finally:
        values.close()
        formulas.close()


def _pptx_shapes(shapes, x_offset=0, y_offset=0, x_scale=1, y_scale=1):
    """Yield shapes with their enclosing-group transformations applied."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        left = x_offset + shape.left * x_scale
        top = y_offset + shape.top * y_scale
        width = shape.width * x_scale
        height = shape.height * y_scale
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            transform = shape._element.grpSpPr.xfrm
            sx = width / transform.chExt.cx if transform.chExt.cx else 1
            sy = height / transform.chExt.cy if transform.chExt.cy else 1
            yield from _pptx_shapes(shape.shapes, left - transform.chOff.x * sx,
                                    top - transform.chOff.y * sy, sx, sy)
        else:
            yield shape, left, top, width, height


def _pptx_to_pdf(src: Path, destination: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    presentation = Presentation(src)
    document = fitz.open()
    emu = 12700  # 914400 EMU/inch divided by 72 pt/inch
    try:
        for slide in presentation.slides:
            page = document.new_page(width=presentation.slide_width / emu, height=presentation.slide_height / emu)
            try:
                background = slide.background.fill.fore_color.rgb
                if background:
                    page.draw_rect(page.rect, color=None, fill=tuple(value / 255 for value in background))
            except (AttributeError, TypeError):
                pass
            for shape, left, top, width, height in _pptx_shapes(slide.shapes):
                rect = fitz.Rect(left / emu, top / emu, (left + width) / emu, (top + height) / emu)
                if rect.is_empty or rect.is_infinite:
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        page.insert_image(rect, stream=shape.image.blob, keep_proportion=False)
                    except Exception as exc:
                        raise UserError("Не удалось прочитать одно из изображений PowerPoint.") from exc
                elif shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        rows.append("<tr>" + "".join(f"<td>{html.escape(cell.text)}</td>" for cell in row.cells) + "</tr>")
                    page.insert_htmlbox(rect, "<table>" + "".join(rows) + "</table>",
                                        css="table{border-collapse:collapse;width:100%}td{border:1px solid #999;padding:3px;font-size:12pt}")
                elif shape.has_text_frame:
                    paragraphs = []
                    for paragraph in shape.text_frame.paragraphs:
                        runs = []
                        for run in paragraph.runs:
                            styles = []
                            if run.font.size:
                                styles.append(f"font-size:{run.font.size.pt}pt")
                            if run.font.bold:
                                styles.append("font-weight:bold")
                            if run.font.italic:
                                styles.append("font-style:italic")
                            try:
                                if run.font.color.rgb:
                                    styles.append(f"color:#{run.font.color.rgb}")
                            except (AttributeError, TypeError):
                                pass
                            runs.append(f'<span style="{";".join(styles)}">{html.escape(run.text)}</span>')
                        paragraphs.append("<p>" + "".join(runs) + "</p>")
                    page.insert_htmlbox(rect, "".join(paragraphs),
                                        css="body{font-family:sans-serif;font-size:18pt}p{margin:0 0 4pt}")
        if not document.page_count:
            raise UserError("Презентация не содержит слайдов.")
        atomic_save(lambda target: document.save(target, garbage=4, deflate=True), destination)
    finally:
        document.close()


def _odt_to_pdf(src: Path, destination: Path) -> None:
    from lxml import etree
    namespaces = {"office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0", "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
                  "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
                  "xlink": "http://www.w3.org/1999/xlink"}
    archive = fitz.Archive()
    with zipfile.ZipFile(src) as source:
        try:
            root = etree.fromstring(source.read("content.xml"), parser=etree.XMLParser(resolve_entities=False, no_network=True))
        except (KeyError, etree.XMLSyntaxError) as exc:
            raise UserError("Документ ODT не содержит читаемого текста.") from exc

        def content(node):
            tag = etree.QName(node).localname
            if tag == "image":
                name = node.get("{" + namespaces["xlink"] + "}href", "")
                if name in source.namelist():
                    archive.add((source.read(name), name))
                    return f'<img src="{html.escape(name, quote=True)}" style="max-width:100%">'
                return ""
            if tag == "s":
                return "&#160;" * min(int(node.get("{" + namespaces["text"] + "}c", "1")), 1000)
            if tag in {"line-break", "tab"}:
                return "<br>" if tag == "line-break" else "&#160;&#160;&#160;&#160;"
            children = html.escape(node.text or "") + "".join(content(child) + html.escape(child.tail or "") for child in node if isinstance(child.tag, str))
            tags = {"p": "p", "h": "h2", "table": "table", "table-row": "tr", "table-cell": "td", "list": "ul", "list-item": "li"}
            if tag in tags:
                output_tag = tags[tag]
                return f"<{output_tag}>{children}</{output_tag}>"
            return children

        body = root.find(".//office:text", namespaces)
        if body is None:
            raise UserError("В ODT-файле не найден текстовый документ.")
        _render_html(content(body), destination, archive=archive)


def office_to_pdf(src: str | Path) -> list[Path]:
    """Basic offline export of DOCX, XLSX, PPTX and ODT with no Office requirement."""
    src = Path(src)
    extension = src.suffix.lower()
    if extension not in OFFICE_EXTENSIONS:
        raise UserError("Поддерживаются DOCX, XLSX, PPTX и ODT. Старые DOC/XLS/PPT сначала сохраните в новом формате.")
    _check_office_archive(src)
    if extension == ".docx":
        return docx_to_pdf(src)
    destination = sibling(src, ext=".pdf")
    exporters = {".xlsx": _xlsx_to_pdf, ".pptx": _pptx_to_pdf, ".odt": _odt_to_pdf}
    try:
        exporters[extension](src, destination)
    except UserError:
        raise
    except Exception as exc:
        raise UserError(f"Не удалось преобразовать документ: {exc}") from exc
    return [destination]


def capture_screen_to_pdf(output_dir: str | Path) -> list[Path]:
    """Run native interactive capture only following the user's explicit action."""
    if sys.platform != "darwin":
        raise UserError("Снимок экрана через системный интерфейс доступен только в macOS.")
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    destination = unique_path(directory / f"Снимок экрана {dt.datetime.now():%Y-%m-%d %H.%M.%S}.pdf")
    with tempfile.TemporaryDirectory(prefix="dna-screen-") as temp:
        screenshot = Path(temp) / "capture.png"
        result = subprocess.run(["/usr/sbin/screencapture", "-i", "-x", "-t", "png", str(screenshot)],
                                capture_output=True, text=True, timeout=300)
        if not screenshot.exists():
            if "permission" in result.stderr.lower() or "could not create image" in result.stderr.lower():
                raise UserError("Разрешите запись экрана для DocuFreedom в настройках macOS → Конфиденциальность и безопасность.")
            return []  # Escape or native capture cancellation.
        with Image.open(screenshot) as image:
            rgb = image.convert("RGB")
            atomic_save(lambda target: rgb.save(target, "PDF", resolution=144), destination)
    return [destination]


def open_scanner() -> None:
    """Open macOS Image Capture. Scanning itself is handled by Apple's device UI."""
    if sys.platform != "darwin":
        raise UserError("Импорт со сканера доступен через «Захват изображений» в macOS.")
    subprocess.run(["/usr/bin/open", "-b", "com.apple.Image_Capture"], check=True, capture_output=True)


def print_pdf(src: str | Path, parent=None) -> bool:
    """Show the native print dialog; call on the Qt GUI thread only."""
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QImage, QPainter
    from PySide6.QtPrintSupport import QPrintDialog, QPrinter
    from PySide6.QtWidgets import QDialog
    with open_pdf(src) as document:
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setResolution(300)
        printer.setDocName(Path(src).stem)
        dialog = QPrintDialog(printer, parent)
        dialog.setMinMax(1, document.page_count)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return False
        painter = QPainter()
        if not painter.begin(printer):
            raise UserError("Не удалось начать печать. Проверьте выбранный принтер.")
        start = max(0, printer.fromPage() - 1)
        end = min(document.page_count, printer.toPage() or document.page_count)
        try:
            for index in range(start, end):
                if index > start and not printer.newPage():
                    raise UserError("Принтер не принял следующую страницу.")
                page = document[index]
                pixmap = page.get_pixmap(dpi=300, alpha=False)
                image = QImage(pixmap.samples, pixmap.width, pixmap.height, pixmap.stride, QImage.Format.Format_RGB888).copy()
                # QPainter's origin is already the printer's printable area.
                target = QRectF(painter.viewport())
                scale = min(target.width() / image.width(), target.height() / image.height())
                width, height = image.width() * scale, image.height() * scale
                rect = QRectF(target.x() + (target.width() - width) / 2, target.y() + (target.height() - height) / 2, width, height)
                painter.drawImage(rect, image)
        finally:
            painter.end()
    return True


def _pdf_to_pptx(src: Path, destination: Path, dpi: int) -> None:
    """One rasterized PDF page per slide, preserving visible page content."""
    from pptx import Presentation
    from pptx.util import Pt
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    with open_pdf(src) as document:
        if not document.page_count:
            raise UserError("PDF не содержит страниц.")
        # PowerPoint permits slide dimensions only from 1 to 56 inches.
        first = document[0].rect
        scale = min(1, 4000 / max(first.width, first.height))
        presentation.slide_width = Pt(max(72, first.width * scale))
        presentation.slide_height = Pt(max(72, first.height * scale))
        for page in document:
            actual_dpi = min(dpi, 72 * math.sqrt(MAX_OCR_PIXELS / (page.rect.width * page.rect.height)))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(actual_dpi / 72, actual_dpi / 72), alpha=False)
            slide = presentation.slides.add_slide(blank)
            factor = min(presentation.slide_width / page.rect.width, presentation.slide_height / page.rect.height)
            width = int(page.rect.width * factor)
            height = int(page.rect.height * factor)
            slide.shapes.add_picture(io.BytesIO(pixmap.tobytes("png")),
                                     int((presentation.slide_width - width) / 2),
                                     int((presentation.slide_height - height) / 2), width=width, height=height)
    atomic_save(lambda target: presentation.save(target), destination)


def _pdf_to_xlsx(src: Path, destination: Path) -> None:
    """Extract detected text tables into worksheets; never execute PDF values as formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    with open_pdf(src) as document:
        for pno, page in enumerate(document):
            tables = page.find_tables().tables
            for tno, table in enumerate(tables):
                data = table.extract()
                if not data:
                    continue
                sheet = workbook.create_sheet(f"Стр {pno + 1} Табл {tno + 1}")
                for rno, row in enumerate(data, 1):
                    for cno, value in enumerate(row, 1):
                        cell = sheet.cell(rno, cno, value)
                        if isinstance(value, str):
                            cell.data_type = "s"
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        if rno == 1:
                            cell.font = Font(bold=True, color="FFFFFF")
                            cell.fill = PatternFill("solid", fgColor="6B2038")
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cno in range(1, sheet.max_column + 1):
                    width = max(len(str(sheet.cell(rno, cno).value or "")) for rno in range(1, min(sheet.max_row + 1, 300)))
                    sheet.column_dimensions[get_column_letter(cno)].width = min(55, max(12, width + 2))
    if not workbook.worksheets:
        raise UserError("В PDF не найдены текстовые таблицы. Для скана сначала выполните OCR; таблицы должны иметь читаемые строки и столбцы.")
    atomic_save(lambda target: workbook.save(target), destination)


def _pdf_to_tiff(src: Path, destination: Path, dpi: int) -> None:
    from PIL.TiffImagePlugin import AppendingTiffWriter

    class TIFFWriter(AppendingTiffWriter):
        # Pillow 13's BufferedIOBase finalizer calls close again after __exit__.
        # Make closing idempotent to avoid finalizing an already closed stream.
        def close(self):
            if not self.f.closed:
                super().close()

    def write(target):
        with open_pdf(src) as document, TIFFWriter(str(target), new=True) as writer:
            for page in document:
                actual_dpi = min(dpi, 72 * math.sqrt(MAX_OCR_PIXELS / (page.rect.width * page.rect.height)))
                pixmap = page.get_pixmap(matrix=fitz.Matrix(actual_dpi / 72, actual_dpi / 72), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                image.save(writer, "TIFF", compression="tiff_lzw", dpi=(actual_dpi, actual_dpi))
                writer.newFrame()
    atomic_save(write, destination)


def export_document(src: str | Path, target: str, *, dpi: int = 150) -> list[Path]:
    """Extra local exports: PDF→PPTX/XLSX/SVG/HTML/TIFF and SVG/EPUB/XPS→PDF."""
    src = Path(src)
    target = target.lower().lstrip(".")
    if not 72 <= dpi <= 600:
        raise UserError("Разрешение должно быть от 72 до 600 dpi.")
    if target == "pdf" and src.suffix.lower() in {".svg", ".epub", ".xps", ".oxps", ".fb2", ".cbz"}:
        destination = sibling(src, ext=".pdf")
        try:
            with fitz.open(src) as source:
                data = source.convert_to_pdf()
            atomic_save(lambda path: Path(path).write_bytes(data), destination)
        except Exception as exc:
            raise UserError(f"Не удалось прочитать документ: {exc}") from exc
        return [destination]
    if src.suffix.lower() != ".pdf":
        raise UserError("Для этого экспорта выберите PDF.")
    if target not in {"pptx", "xlsx", "svg", "html", "tiff", "tif"}:
        raise UserError("Неподдерживаемый формат экспорта.")
    destination = sibling(src, ext="." + target)
    if target == "pptx":
        _pdf_to_pptx(src, destination, dpi)
    elif target == "xlsx":
        _pdf_to_xlsx(src, destination)
    elif target in {"tif", "tiff"}:
        _pdf_to_tiff(src, destination, dpi)
    elif target == "html":
        with open_pdf(src) as document:
            pages = []
            for index, page in enumerate(document):
                # Embedded base64 images and absolute text placement, no external assets.
                pages.append(page.get_text("html").replace('id="page0"', f'id="page{index}"', 1))
        result = ('<!doctype html><html><head><meta charset="utf-8"><title>' + html.escape(src.stem) +
                  '</title><style>body{background:#ddd;margin:24px}body>div{position:relative;margin:0 auto 24px;background:white;overflow:hidden}</style>' +
                  '</head><body>' + "\n".join(pages) + '</body></html>')
        atomic_save(lambda path: Path(path).write_text(result, encoding="utf-8"), destination)
    elif target == "svg":
        outputs = []
        with open_pdf(src) as document:
            for index, page in enumerate(document):
                path = sibling(src, f"Страница {index + 1}" if document.page_count > 1 else "", ".svg")
                data = page.get_svg_image(text_as_path=True)
                atomic_save(lambda output: Path(output).write_text(data, encoding="utf-8"), path)
                outputs.append(path)
        return outputs
    return [destination]


def install_pdf_service(app_path: str | Path = "/Applications/DocuFreedom.app") -> Path:
    """User-invoked setup for macOS Print → PDF → DocuFreedom; no system printer/driver."""
    if sys.platform != "darwin":
        raise UserError("Меню PDF печати доступно только в macOS.")
    application = Path(app_path)
    if not application.is_dir() or application.suffix != ".app":
        raise UserError("Сначала установите DocuFreedom в папку «Программы».")
    services = Path.home() / "Library" / "PDF Services"
    services.mkdir(parents=True, exist_ok=True)
    destination = services / "Открыть в DocuFreedom.app"
    if destination.is_symlink() and destination.resolve() == application.resolve():
        return destination
    if destination.exists() or destination.is_symlink():
        raise UserError("Элемент «Открыть в DocuFreedom» уже существует в меню PDF. Удалите старый элемент вручную.")
    destination.symlink_to(application, target_is_directory=True)
    return destination
