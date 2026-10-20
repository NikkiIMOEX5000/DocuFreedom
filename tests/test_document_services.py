"""Content and real on-device OCR checks for local document services."""
import io
import sys
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from unittest.mock import patch

import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw, ImageFont

from dragshift.document_services import (
    OCRLine, UserError, capture_screen_to_pdf, export_document, html_to_pdf,
    install_pdf_service, ocr_document, office_to_pdf, recognize_image, webpage_to_pdf,
)


@pytest.fixture
def source_pdf(tmp_path):
    path = tmp_path / "source.pdf"
    document = fitz.open()
    for index in range(2):
        page = document.new_page(width=400, height=300)
        page.insert_text((30, 40), f"DNA.apps page {index + 1}", fontsize=20)
        page.draw_rect(fitz.Rect(30, 70, 180, 150), fill=(0.4, 0.1, 0.2))
    document.save(path)
    document.close()
    return path


@pytest.fixture
def readable_image(tmp_path):
    font_path = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
    if not font_path.exists():
        pytest.skip("Test fixture requires the macOS Arial font")
    image = Image.new("RGB", (1500, 500), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 60)
    draw.text((80, 80), "DNA.apps 2026", font=font, fill="black")
    draw.text((80, 190), "Проверка документа", font=font, fill="black")
    path = tmp_path / "ocr-input.png"
    image.save(path)
    return path


@pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision requires macOS")
def test_actual_vision_english_and_russian(readable_image):
    pytest.importorskip("Vision")
    lines = recognize_image(Image.open(readable_image))
    text = "\n".join(line.text for line in lines)
    assert "DNA.apps 2026" in text
    assert "Проверка документа" in text
    assert all(0 <= line.x <= 1 and 0 <= line.y <= 1 for line in lines)


@pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision requires macOS")
def test_real_searchable_pdf_preserves_appearance_and_locations(readable_image):
    pytest.importorskip("Vision")
    progress = []
    result = ocr_document(readable_image, progress=lambda *args: progress.append(args))[0]
    with fitz.open(result) as document:
        page = document[0]
        assert "Проверка документа" in page.get_text()
        english = page.search_for("DNA.apps")[0]
        russian = page.search_for("Проверка")[0]
        assert english.y0 < russian.y0
        assert len(page.get_images()) == 1
        pixmap = page.get_pixmap(dpi=200)
        rendered = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        original = Image.open(readable_image)
        assert rendered.size == original.size
        # The searchable layer is invisible; exact original page pixels remain.
        assert rendered.tobytes() == original.tobytes()
    assert progress == [(1, 1)]
    assert readable_image.exists()


def test_ocr_text_multipage_and_blank_failure(tmp_path):
    source = tmp_path / "scan.tiff"
    first = Image.new("RGB", (200, 100), "white")
    first.save(source, save_all=True, append_images=[first.copy()])
    line = OCRLine("Пример", 0.1, 0.1, 0.6, 0.2, 0.9)
    with patch("dragshift.document_services.recognize_image", return_value=[line]):
        result = ocr_document(source, output="txt")[0]
    assert result.read_text(encoding="utf-8") == "Пример\n\f\nПример"
    with patch("dragshift.document_services.recognize_image", return_value=[]), pytest.raises(UserError, match="Текст не найден"):
        ocr_document(source)
    assert not (tmp_path / "scan OCR.pdf").exists()


def test_ocr_pdf_input_and_validation(source_pdf):
    line = OCRLine("Sample", 0.1, 0.1, 0.5, 0.1, 1)
    with patch("dragshift.document_services.recognize_image", return_value=[line]):
        result = ocr_document(source_pdf)[0]
    with fitz.open(result) as document:
        assert document.page_count == 2
        assert document[0].rect == fitz.Rect(0, 0, 400, 300)
        assert "Sample" in document[1].get_text()
    with pytest.raises(UserError, match="PDF или TXT"):
        ocr_document(source_pdf, output="docx")
    with pytest.raises(UserError, match="Разрешение"):
        ocr_document(source_pdf, dpi=0)


def test_html_unicode_images_links_no_scripts(tmp_path):
    source = tmp_path / "local.html"
    Image.new("RGB", (30, 30), "red").save(tmp_path / "image.png")
    source.write_text('<html><meta charset="utf-8"><h1>Пример HTML</h1><p>Text 123</p><img src="image.png"><a href="https://example.com">Example link</a><script>secret script text</script><img src="https://example.com/remote.png"></html>', encoding="utf-8")
    result = html_to_pdf(source)[0]
    with fitz.open(result) as document:
        assert "Пример HTML" in document[0].get_text()
        assert "secret script" not in document[0].get_text()
        assert len(document[0].get_images()) == 1
        assert document[0].get_links()[0]["uri"] == "https://example.com"
    assert source.exists()


def test_html_images_cannot_escape_source_folder(tmp_path):
    folder = tmp_path / "sub"
    folder.mkdir()
    Image.new("RGB", (30, 30), "red").save(tmp_path / "private.png")
    source = folder / "local.html"
    source.write_text('<p>Allowed text</p><img src="../private.png">', encoding="utf-8")
    result = html_to_pdf(source)[0]
    with fitz.open(result) as document:
        assert len(document[0].get_images()) == 0
        assert "Allowed text" in document[0].get_text()


def test_webpage_explicit_fetch_does_not_fetch_images(tmp_path):
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write('<h1>Страница</h1><p>Requested page</p><img src="/tracking.png">'.encode("utf-8"))
        def log_message(self, *args):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        result = webpage_to_pdf(f"http://127.0.0.1:{server.server_port}/requested", tmp_path)[0]
    finally:
        server.shutdown()
        server.server_close()
        thread.join()
    assert requests == ["/requested"]
    with fitz.open(result) as document:
        assert "Страница" in document[0].get_text()
    for bad in ["file:///etc/passwd", "https://user:pass@example.com", "example.com"]:
        with pytest.raises(UserError):
            webpage_to_pdf(bad, tmp_path)


def test_xlsx_preserves_sheets_unicode_and_calculated_results(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"
    sheet.append(["Товар", "Стоимость"])
    sheet.append(["Бумага", 42.5])
    sheet.append(["Формула", "=B2*2"])
    hidden = workbook.create_sheet("hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["HIDDEN_SECRET"])
    workbook.create_sheet("Second").append(["Second sheet text"])
    source = tmp_path / "book.xlsx"
    workbook.save(source)
    result = office_to_pdf(source)[0]
    with fitz.open(result) as document:
        text = "\n".join(page.get_text() for page in document)
        assert document.page_count == 2
        assert "Бумага" in text and ("42.5" in text or "42,5" in text)
        assert "85" in text
        assert "=B2*2" not in text
        assert "Second sheet text" in text
        assert "HIDDEN_SECRET" not in text


def test_pptx_text_images_and_table(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(1))
    textbox.text_frame.text = "Презентация DNA.apps"
    textbox.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    image = io.BytesIO()
    Image.new("RGB", (100, 100), "red").save(image, "PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(0.5), Inches(2), Inches(1), Inches(1))
    table = slide.shapes.add_table(2, 2, Inches(2), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Table value"
    source = tmp_path / "slides.pptx"
    presentation.save(source)
    result = office_to_pdf(source)[0]
    with fitz.open(result) as document:
        assert document.page_count == 1
        assert "Презентация DNA.apps" in document[0].get_text()
        assert "Table value" in document[0].get_text()
        assert len(document[0].get_images()) == 1


def test_odt_paragraph_table_image(tmp_path):
    image = io.BytesIO()
    Image.new("RGB", (30, 30), "red").save(image, "PNG")
    source = tmp_path / "document.odt"
    content = '''<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content office:version="1.2"
        xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
        xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
        xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
        xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
        xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
        xmlns:xlink="http://www.w3.org/1999/xlink">
      <office:body><office:text>
        <text:h text:outline-level="1">Документ ODT</text:h><text:p>Hello ODT</text:p>
        <table:table table:name="Example"><table:table-column/>
          <table:table-row><table:table-cell office:value-type="string">
            <text:p>Cell text</text:p>
          </table:table-cell></table:table-row>
        </table:table>
        <text:p><draw:frame draw:name="Illustration" text:anchor-type="as-char"
            svg:width="2cm" svg:height="2cm">
          <draw:image xlink:href="Pictures/red.png" xlink:type="simple"
              xlink:show="embed" xlink:actuate="onLoad"/>
        </draw:frame></text:p>
      </office:text></office:body>
    </office:document-content>'''
    manifest = '''<?xml version="1.0" encoding="UTF-8"?>
    <manifest:manifest manifest:version="1.2"
        xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
      <manifest:file-entry manifest:full-path="/"
          manifest:media-type="application/vnd.oasis.opendocument.text"/>
      <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
      <manifest:file-entry manifest:full-path="Pictures/red.png" manifest:media-type="image/png"/>
    </manifest:manifest>'''
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("META-INF/manifest.xml", manifest)
        archive.writestr("content.xml", content)
        archive.writestr("Pictures/red.png", image.getvalue())
    result = office_to_pdf(source)[0]
    with fitz.open(result) as document:
        text = document[0].get_text()
        assert "Документ ODT" in text and "Cell text" in text
        assert len(document[0].get_images()) == 1


def test_docx_and_unsupported_office(tmp_path):
    from docx import Document
    source = tmp_path / "document.docx"
    document = Document()
    document.add_paragraph("Docx content")
    document.save(source)
    result = office_to_pdf(source)[0]
    with fitz.open(result) as pdf:
        assert "Docx content" in pdf[0].get_text()
    with pytest.raises(UserError, match="Старые"):
        office_to_pdf(tmp_path / "old.doc")
    damaged = tmp_path / "broken.xlsx"
    damaged.write_text("not a zip", encoding="utf-8")
    with pytest.raises(UserError, match="Повреждённый"):
        office_to_pdf(damaged)


def test_pdf_to_pptx_retains_each_page(source_pdf):
    from pptx import Presentation
    result = export_document(source_pdf, "PPTX")[0]
    presentation = Presentation(result)
    assert len(presentation.slides) == 2
    assert all(len(slide.shapes) == 1 for slide in presentation.slides)
    first_image = Image.open(io.BytesIO(presentation.slides[0].shapes[0].image.blob))
    assert first_image.width > first_image.height


def test_pdf_to_xlsx_extracts_table_as_data(tmp_path):
    from openpyxl import load_workbook
    source = tmp_path / "table.pdf"
    document = fitz.open()
    page = document.new_page(width=400, height=400)
    for x in (20, 160, 300):
        page.draw_line((x, 30), (x, 150))
    for y in (30, 70, 110, 150):
        page.draw_line((20, y), (300, y))
    for row, values in enumerate([["Name", "Value"], ["One", "42"], ["Formula", "=1+1"]]):
        for col, value in enumerate(values):
            page.insert_text((30 + 140 * col, 55 + row * 40), value)
    document.save(source)
    document.close()
    result = export_document(source, "XLSX")[0]
    workbook = load_workbook(result)
    assert workbook.active["A2"].value == "One"
    assert workbook.active["B3"].value == "=1+1"
    assert workbook.active["B3"].data_type == "s"
    workbook.close()


def test_pdf_to_xlsx_without_tables_has_actionable_error(source_pdf):
    with pytest.raises(UserError, match="не найдены текстовые таблицы"):
        export_document(source_pdf, "xlsx")


def test_pdf_tiff_svg_html_exports(source_pdf):
    tiff = export_document(source_pdf, "TIFF")[0]
    with Image.open(tiff) as image:
        assert image.n_frames == 2
        image.seek(1)
        assert image.getpixel((100, 220))[0] > 50
    svgs = export_document(source_pdf, "SVG")
    assert len(svgs) == 2
    assert all("<svg" in path.read_text(encoding="utf-8") for path in svgs)
    markup = export_document(source_pdf, "HTML")[0].read_text(encoding="utf-8")
    assert "DNA.apps page 1" in markup and "DNA.apps page 2" in markup
    assert 'id="page0"' in markup and 'id="page1"' in markup


def test_svg_to_pdf(tmp_path):
    source = tmp_path / "vector.svg"
    source.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200"><rect width="400" height="200" fill="#67213b"/><text x="30" y="100" fill="white">Vector document</text></svg>', encoding="utf-8")
    result = export_document(source, "PDF")[0]
    with fitz.open(result) as document:
        assert document.page_count == 1
        assert document[0].get_drawings()
        assert "Vector document" in document[0].get_text()


@pytest.mark.skipif(sys.platform != "darwin", reason="macOS Print PDF service")
def test_pdf_service_idempotent_does_not_replace(tmp_path):
    application = tmp_path / "DNA.apps.app"
    application.mkdir()
    with patch("pathlib.Path.home", return_value=tmp_path):
        service = install_pdf_service(application)
        assert service.is_symlink()
        assert install_pdf_service(application) == service
        service.unlink()
        service.write_text("user content", encoding="utf-8")
        with pytest.raises(UserError, match="уже существует"):
            install_pdf_service(application)
        assert service.read_text(encoding="utf-8") == "user content"


@pytest.mark.skipif(sys.platform != "darwin", reason="macOS interactive capture backend")
def test_capture_cancellation_does_not_create_pdf(tmp_path):
    with patch("dragshift.document_services.subprocess.run") as run:
        run.return_value.stderr = ""
        assert capture_screen_to_pdf(tmp_path) == []
        args = run.call_args.args[0]
        assert args[:3] == ["/usr/sbin/screencapture", "-i", "-x"]
    assert list(tmp_path.glob("*.pdf")) == []


def test_epub_to_pdf(tmp_path):
    source = tmp_path / "book.epub"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", '''<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>''')
        archive.writestr("OEBPS/content.opf", '''<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="book"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Test book</dc:title><dc:language>en</dc:language><dc:identifier id="book">test</dc:identifier></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chapter"/></spine></package>''')
        archive.writestr("OEBPS/chapter.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Book</title></head><body><h1>EPUB chapter</h1><p>Text from a book</p></body></html>')
    result = export_document(source, "PDF")[0]
    with fitz.open(result) as document:
        assert "EPUB chapter" in document[0].get_text()
        assert "Text from a book" in document[0].get_text()


@pytest.mark.skipif(sys.platform != "darwin", reason="macOS interactive capture backend")
def test_capture_success_converts_selected_image_only(tmp_path):
    def screenshot(args, **kwargs):
        Image.new("RGB", (300, 200), (120, 30, 60)).save(args[-1])
        return type("Result", (), {"stderr": ""})()
    with patch("dragshift.document_services.subprocess.run", side_effect=screenshot):
        result = capture_screen_to_pdf(tmp_path)[0]
    with fitz.open(result) as document:
        assert document.page_count == 1
        assert document[0].rect.width == 150
        pixmap = document[0].get_pixmap()
        assert abs(pixmap.pixel(10, 10)[0] - 120) <= 2


def test_native_print_pipeline_without_hardware(source_pdf, tmp_path):
    from PySide6.QtWidgets import QApplication, QDialog
    from PySide6.QtPrintSupport import QPrinter
    from dragshift.document_services import print_pdf
    app = QApplication.instance() or QApplication([])
    destination = tmp_path / "printed.pdf"

    class Dialog:
        def __init__(self, printer, parent=None):
            printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
            printer.setOutputFileName(str(destination))
        def setMinMax(self, *args):
            pass
        def exec(self):
            return QDialog.DialogCode.Accepted

    with patch("PySide6.QtPrintSupport.QPrintDialog", Dialog):
        assert print_pdf(source_pdf)
    with fitz.open(destination) as document:
        assert document.page_count == 2
        assert all(page.get_images() for page in document)
    assert app is not None


def test_transparent_ocr_input_is_composited_on_white(tmp_path):
    path = tmp_path / "transparent.png"
    Image.new("RGBA", (200, 100), (0, 0, 0, 0)).save(path)
    captured = []
    def recognition(image, languages):
        captured.append(image.getpixel((10, 10)))
        return [OCRLine("Sample", 0.1, 0.1, 0.5, 0.1, 1)]
    with patch("dragshift.document_services.recognize_image", side_effect=recognition):
        ocr_document(path)
    assert captured == [(255, 255, 255)]


def test_html_utf8_without_charset_declaration(tmp_path):
    source = tmp_path / "unicode.html"
    source.write_text("<h1>Русский без метаданных</h1>", encoding="utf-8")
    result = html_to_pdf(source)[0]
    with fitz.open(result) as document:
        assert "Русский без метаданных" in document[0].get_text()
