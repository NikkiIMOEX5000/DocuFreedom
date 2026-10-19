"""Synthetic Word layout fixture; contains no user files or personal information.

Generate a document for a manual Word/Preview comparison with:
    python tests/office_fidelity_fixture.py /tmp/office-fidelity.docx
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw


def _picture(color: str, label: str) -> io.BytesIO:
    image = Image.new("RGB", (120, 80), color)
    draw = ImageDraw.Draw(image)
    draw.rectangle((3, 3, 116, 76), outline="white", width=3)
    draw.text((12, 32), label, fill="white")
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    return stream


def _floating_picture(paragraph) -> None:
    """Anchor a picture to absolute page coordinates, not paragraph flow."""
    shape = paragraph.add_run().add_picture(_picture("#1256b0", "FLOAT"), width=Inches(1))
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "251658240", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    anchor.append(simple_position)
    for axis, points in (("H", 420), ("V", 100)):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = str(int(Pt(points)))
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    effect_extent = inline.find(qn("wp:effectExtent"))
    if effect_extent is not None:
        anchor.append(effect_extent)
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)
    inline.getparent().replace(inline, anchor)


def _numbering(document) -> str:
    root = document.part.numbering_part.element
    abstract_id = max(int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))) + 1
    num_id = max(int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1.")):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        level.append(element)
    abstract.append(level)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    num.append(reference)
    root.append(num)
    return str(num_id)


def create_fidelity_docx(destination: Path) -> Path:
    document = Document()
    document.core_properties.author = "DocuFreedom regression suite"
    document.core_properties.title = "Office layout fidelity fixture"
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1
    body_style = document.styles.add_style("Fixture body", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = normal
    body_style.font.name = "Times New Roman"
    body_style.font.size = Pt(12)

    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.75), Inches(.75)
    section.left_margin, section.right_margin = Inches(1), Inches(.75)
    section.header_distance, section.footer_distance = Inches(.3), Inches(.3)
    section.header.paragraphs[0].text = "PORTRAIT_HEADER"
    section.footer.paragraphs[0].text = "PORTRAIT_FOOTER"

    title = document.add_paragraph("Office layout fidelity fixture", "Title")
    title.runs[0].font.name = "Times New Roman"
    title.runs[0].font.size = Pt(22)
    title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    # python-docx's stock Title style contains a decorative bottom border.
    for element in document.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        element.getparent().remove(element)
    document.add_paragraph("POSITION_PROBE", "Fixture body")
    document.add_paragraph("Кириллица: Привет, мир! Ёжик и таблица.", "Fixture body")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("STYLE_PROBE")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.bold, run.italic = True, True
    run.font.color.rgb = RGBColor.from_string("1256B0")
    centered = document.add_paragraph("CENTERED_PROBE")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    number_id = _numbering(document)
    for text in ("First numbered step", "Second numbered step"):
        paragraph = document.add_paragraph(text)
        props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        props.get_or_add_ilvl().val = 0
        props.get_or_add_numId().val = int(number_id)

    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    for column in table.columns:
        column.width = Inches(2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "HORIZONTAL_MERGE"
    table.cell(0, 2).text = "LAST_COLUMN"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "VERTICAL_MERGE"
    table.cell(1, 1).text = "CELL_B2"
    table.cell(1, 2).text = "CELL_C2"
    table.cell(2, 1).text = "CELL_B3"
    table.cell(2, 2).text = "CELL_C3"
    # Leave optional wp:inline distances absent, as python-docx normally does.
    # They must not introduce extra padding for an inline drawing.
    document.add_picture(_picture("#70283d", "INLINE"), width=Inches(1.2))
    _floating_picture(document.add_paragraph())

    document.add_page_break()
    document.add_paragraph("EXPLICIT_BREAK_PAGE", "Heading 1")
    document.add_paragraph("This content must begin on the second portrait page.")

    landscape = document.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = Mm(297), Mm(210)
    landscape.left_margin, landscape.right_margin = Inches(.5), Inches(.5)
    landscape.top_margin, landscape.bottom_margin = Inches(.6), Inches(.6)
    landscape.header.is_linked_to_previous = False
    landscape.footer.is_linked_to_previous = False
    landscape.header.paragraphs[0].text = "LANDSCAPE_HEADER"
    landscape.footer.paragraphs[0].text = "LANDSCAPE_FOOTER"
    document.add_paragraph("LANDSCAPE_SECTION", "Heading 1")
    document.add_paragraph("Landscape content has a different page size and margins.")
    document.add_paragraph("LANDSCAPE_MARGIN_PROBE")
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    print(create_fidelity_docx(parser.parse_args().output))
