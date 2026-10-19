"""Local browsing, bounded previews, editor handoffs and safe text copy saves."""
from __future__ import annotations

import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from pathlib import Path
from types import SimpleNamespace
import zipfile

import pytest
from PIL import Image
from PySide6.QtCore import QSettings, Qt
from PySide6.QtWidgets import QApplication, QMessageBox

from dragshift.convert import fitz
from dragshift.library_workspace import LibraryWorkspace, PREPARE_EDIT_ACTION, office_text_preview
from dragshift.report import Result
from dragshift.util import UserError


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


class Controller:
    def __init__(self):
        self.jobs = []
        self.windows = []
        self.results = []
        self.items = []
    def submit(self, action, sources, fn, **kwargs):
        self.jobs.append((action, sources, fn, kwargs))
    def open_window(self, window):
        self.windows.append(window)
    def notify(self, *args):
        self.results.append(args)
    def run_item(self, item, paths):
        self.items.append((item, paths))


@pytest.fixture
def files(tmp_path):
    (tmp_path / "nested").mkdir()
    (tmp_path / "nested" / "do-not-index.txt").write_text("not scanned")
    text = tmp_path / "Пример.txt"
    text.write_text("Первоначальный текст", encoding="utf-8")
    (tmp_path / ".hidden.txt").write_text("hidden")
    image = tmp_path / "image.png"
    Image.new("RGB", (640, 400), "#f3e8ec").save(image)
    pdf = tmp_path / "document.pdf"
    with fitz.open() as doc:
        for number in range(3):
            page = doc.new_page(width=300, height=400)
            page.insert_text((25, 40), f"Page {number + 1}")
        doc.save(pdf)
    office = tmp_path / "document.docx"
    from docx import Document
    doc = Document()
    doc.add_heading("Заголовок Word", 0)
    doc.add_paragraph("Текстовый просмотр без изменения оригинала.")
    doc.save(office)
    return SimpleNamespace(directory=tmp_path, text=text, image=image, pdf=pdf, office=office)


@pytest.fixture
def workspace(qapp):
    widget = LibraryWorkspace(Controller())
    yield widget
    widget.text.document().setModified(False)
    widget._close_pdf()
    for window in widget.controller.windows:
        if hasattr(window, "dirty"):
            window.dirty = False
        window.close()
    widget.close()
    widget.deleteLater()
    qapp.processEvents()


def item_names(workspace):
    return [workspace.files.item(index).text() for index in range(workspace.files.count())]


def test_folder_browser_only_lists_selected_directory_and_persists(qapp, files, tmp_path):
    store = QSettings(str(tmp_path / "preferences.ini"), QSettings.IniFormat)
    workspace = LibraryWorkspace(Controller(), store=store)
    assert workspace.folder is None
    assert workspace.set_folder(files.directory)
    assert "nested" in item_names(workspace)
    assert "do-not-index.txt" not in item_names(workspace)
    assert ".hidden.txt" not in item_names(workspace)
    assert store.value("library/last_folder") == str(files.directory.resolve())
    workspace.hidden.setChecked(True)
    assert ".hidden.txt" in item_names(workspace)
    workspace.search.setText("document")
    visible = [workspace.files.item(index).text() for index in range(workspace.files.count()) if not workspace.files.item(index).isHidden()]
    assert visible == ["document.docx", "document.pdf"]
    workspace.close()
    restored = LibraryWorkspace(Controller(), store=store)
    assert restored.folder == files.directory.resolve()
    restored.close()


def test_pdf_preview_page_controls_pause_and_resume(workspace, files, monkeypatch):
    original = files.pdf.read_bytes()
    assert workspace.open_paths([files.pdf])
    assert workspace._preview_kind == "pdf"
    assert workspace.page.maximum() == 3
    workspace.next_button.click()
    assert workspace.page.value() == 2
    assert workspace.picture.pixmap().width() <= 1800
    workspace.set_processing(True)
    assert workspace._pdf is None
    assert not workspace.preview_timer.isActive()
    assert not workspace.edit_button.isEnabled()
    with monkeypatch.context() as patch:
        patch.setattr(fitz, "open", lambda *_a, **_k: pytest.fail("No PDF opens during a worker job"))
        workspace.render_preview()
        workspace.preview_file(files.pdf)
        workspace.edit_selected()
    workspace.set_processing(False)
    assert workspace._pdf is not None
    assert workspace.edit_button.isEnabled()
    assert files.pdf.read_bytes() == original


def test_image_and_office_previews_are_read_only_and_actions_work(workspace, files):
    image_original = files.image.read_bytes()
    office_original = files.office.read_bytes()
    workspace.open_paths([files.image])
    assert workspace._preview_kind == "image"
    workspace.edit_button.click()
    assert workspace.controller.items[0][0].key == "tool:Edit"
    workspace.open_paths([files.office])
    assert workspace._preview_kind == "office"
    assert workspace.text.isReadOnly()
    assert "Заголовок Word" in workspace.text.toPlainText()
    workspace.edit_button.click()
    assert workspace.controller.jobs[0][0] == f"library:{PREPARE_EDIT_ACTION}"
    assert files.image.read_bytes() == image_original
    assert files.office.read_bytes() == office_original


def test_text_editor_saves_copy_never_original(workspace, files, tmp_path):
    original = files.text.read_bytes()
    workspace.open_paths([files.text])
    assert workspace.text.isReadOnly()
    workspace.edit_button.click()
    assert not workspace.text.isReadOnly()
    workspace.text.insertPlainText("Новый текст\n")
    output = workspace.write_text_copy(tmp_path / "edited.txt")
    assert "Новый текст" in output.read_text(encoding="utf-8")
    assert files.text.read_bytes() == original
    with pytest.raises(UserError, match="исходный"):
        workspace.write_text_copy(files.text)


def test_unsaved_text_cancel_preserves_document_selection(workspace, files, monkeypatch):
    workspace.open_paths([files.text])
    workspace.edit_selected()
    workspace.text.insertPlainText("unsaved")
    monkeypatch.setattr(QMessageBox, "question", lambda *_: QMessageBox.Cancel)
    assert not workspace.set_folder(files.directory / "nested")
    assert workspace.current_path == files.text.resolve()
    assert "unsaved" in workspace.text.toPlainText()


def test_large_text_preview_cannot_save_truncated_content(workspace, tmp_path):
    source = tmp_path / "large.txt"
    source.write_text("x" * 300000)
    workspace.open_paths([source])
    assert not workspace._text_complete
    assert not workspace.edit_button.isEnabled()
    with pytest.raises(UserError):
        workspace.write_text_copy(tmp_path / "copy.txt")
    assert source.stat().st_size == 300000


def test_editor_handoff_and_tool_sources(workspace, files):
    workspace.open_paths([files.pdf])
    workspace.edit_selected()
    assert len(workspace.controller.windows) == 1
    assert workspace.controller.windows[0].source == files.pdf.resolve()
    received = []
    workspace.tools_requested.connect(lambda paths: received.append(paths))
    workspace.open_tools()
    assert received == [[str(files.pdf.resolve())]]


def test_office_conversion_opens_generated_pdf_editor(workspace, files, tmp_path):
    workspace.open_paths([files.office])
    workspace.edit_selected()
    _, source, convert, _ = workspace.controller.jobs[0]
    outputs = convert()
    assert outputs and outputs[0].is_file()
    workspace.operation_finished(Result(PREPARE_EDIT_ACTION, "", outputs, [Path(source[0])]))
    assert workspace.current_path == outputs[0].resolve()
    assert len(workspace.controller.windows) == 1
    assert workspace.controller.windows[0].source == outputs[0].resolve()


def test_office_preview_rejects_excessive_expanded_archive(tmp_path):
    bad = tmp_path / "large.docx"
    with zipfile.ZipFile(bad, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (17 * 1024 * 1024))
    with pytest.raises(UserError, match="велико"):
        office_text_preview(bad)


def test_password_unlock_and_page_position_survive_background_work(workspace, files, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QInputDialog
    from dragshift.pdf_ops import encrypt_pdf
    protected = encrypt_pdf(files.pdf, password="secret")[0]
    workspace.open_paths([protected])
    assert workspace._pdf is None
    assert not workspace.unlock_button.isHidden()
    monkeypatch.setattr(QInputDialog, "getText", lambda *_: ("secret", True))
    workspace.unlock_pdf()
    assert workspace._pdf is not None
    workspace.page.setValue(3)
    workspace.set_processing(True)
    workspace.set_processing(False)
    assert workspace.page.value() == 3
    assert workspace._pdf is not None
    assert "приостановлен" not in workspace.preview_info.text()


def test_desktop_starts_with_library_and_routes_selection_to_studio(qapp, files):
    from dragshift.desktop import DesktopWindow
    controller = Controller()
    controller.settings = SimpleNamespace(get=lambda key: "next")
    controller.trigger_label = lambda: "Shift ⇧"
    controller.open_dialog = lambda: None
    controller.open_settings = controller.show_welcome = controller.request_quit = lambda: None
    controller.build_rings = lambda paths: {}
    window = DesktopWindow(controller)
    assert window.pages.currentIndex() == 3 and window.library_nav.isChecked()
    window.preview_file(files.pdf)
    assert window.library.current_path == files.pdf.resolve()
    window.open_pdf_studio()
    assert window.active_files() == [str(files.pdf.resolve())]
    window.open_library(files.office)
    window.library.open_convert()
    assert window.pages.currentIndex() == 2  # Office uses its dedicated PDF Studio converter.
    assert window.active_files() == [str(files.office.resolve())]
    window.library._close_pdf()
    window.close()
